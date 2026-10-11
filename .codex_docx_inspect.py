from docx import Document
from pathlib import Path
import sys
sys.stdout.reconfigure(encoding="utf-8")

p = Path(r"Bao_Cao/NihongoLife_Bao_Cao_Capstone_v4.docx")
d = Document(p)
print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "sections", len(d.sections))
for i, para in enumerate(d.paragraphs):
    t = " ".join(para.text.split())
    if t and any(k in t.lower() for k in ("hud", "video", "trailer", "giao diện", "chỉ số", "player")):
        print(f"P{i}\t{para.style.name}\t{t[:500]}")
for ti, table in enumerate(d.tables):
    text = " | ".join(" ".join(c.text.split()) for row in table.rows for c in row.cells)
    if any(k in text.lower() for k in ("hud", "video", "trailer", "giao diện", "chỉ số", "player")):
        print(f"T{ti}\t{text[:1200]}")
print("--- RANGE 200-330 ---")
for i in range(200, min(331, len(d.paragraphs))):
    t = " ".join(d.paragraphs[i].text.split())
    if t:
        print(f"P{i}\t{d.paragraphs[i].style.name}\t{t[:800]}")
