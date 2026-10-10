"""Builds the trailer script and the gameplay-demo script (DOCX) from the real recordings: timecodes come from each
recording's timeline.tsv and every row carries a thumbnail grabbed from the recorded frames.
Usage: python build_scripts_v2.py <trailer_frames_dir> <demo_frames_dir>"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).parent
OUT = HERE.parent
TR, DE = Path(sys.argv[1]), Path(sys.argv[2])
THUMBS = HERE / "thumbs"; THUMBS.mkdir(exist_ok=True)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def timeline(folder):
    rows = [l.split("\t") for l in (folder / "timeline.tsv").read_text().split("\n") if l.strip()]
    total = len(list(folder.glob("f*.jpg")))
    out = []
    for k, (f, i) in enumerate(rows):
        end = int(rows[k + 1][0]) if k + 1 < len(rows) else total
        out.append((i, int(f), end))
    return out, total


def tc(frame):
    s = frame / 30.0
    return f"{int(s // 60):02d}:{s % 60:04.1f}"


def thumb(folder, start, end, name, frac=0.55):
    f = int(start + (end - start) * frac)
    p = THUMBS / f"{name}.jpg"
    Image.open(folder / f"f{f:05d}.jpg").resize((480, 270), Image.LANCZOS).save(p, quality=85)
    return p


# ─────────── document helpers ───────────

def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(k), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "Yu Gothic")
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(k)) is not None:
            del fonts.attrib[qn(k)]
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color


def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.2
    for lvl, size in ((1, 15), (2, 13)):
        h = d.styles[f"Heading {lvl}"]
        h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = NAVY
        rf = h.element.rPr.rFonts
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
        rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman"); rf.set(qn("w:eastAsia"), "Yu Gothic")
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = Cm(2.5); s.right_margin = Cm(2)
    footer(s)
    return d


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); set_font(r, 10)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._element.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text; r._element.append(e)


def landscape(d):
    s = d.add_section()
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21)
    s.left_margin = s.right_margin = Cm(1.6); s.top_margin = s.bottom_margin = Cm(1.6)
    return s


def portrait(d):
    s = d.add_section()
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = Cm(2.5); s.right_margin = Cm(2)


def title_block(d, kicker, title, sub):
    for text, size, bold, color in ((kicker, 12, False, RGBColor(0x55, 0x60, 0x70)), (title, 22, True, NAVY), (sub, 12.5, False, None)):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size, bold, color)
    rule(d)


def rule(d):
    p = d.add_paragraph()
    pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "1F3A5F")): b.set(qn(k), v)
    pbdr.append(b); p._p.get_or_add_pPr().append(pbdr)


def para(d, text, bold=False, italic=False, size=12, align=None):
    p = d.add_paragraph()
    if align == "j": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # **bold** spans
    parts = text.split("**")
    for k, part in enumerate(parts):
        if part: set_font(p.add_run(part), size, bold or k % 2 == 1, italic=italic)
    return p


def bullets(d, items, size=12):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        parts = it.split("**")
        for k, part in enumerate(parts):
            if part: set_font(p.add_run(part), size, k % 2 == 1)


def shade(cell, hex_):
    tcpr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_); tcpr.append(s)


def table(d, headers, rows, widths, size=10, header_fill="1F3A5F", images=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for k, h in enumerate(headers):
        c = t.rows[0].cells[k]; c.text = ""
        set_font(c.paragraphs[0].add_run(h), size, True, RGBColor(0xFF, 0xFF, 0xFF)); shade(c, header_fill)
    trpr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trpr.append(cs)
        for k, val in enumerate(row):
            c = cells[k]; c.text = ""
            if images and k in images and val:
                c.paragraphs[0].add_run().add_picture(str(val), width=widths[k] - Cm(0.25))
                continue
            lines = str(val).split("\n")
            for li, line in enumerate(lines):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.05
                parts = line.split("**")
                for j, part in enumerate(parts):
                    if part: set_font(p.add_run(part), size, j % 2 == 1)
            if r_i % 2 == 1: shade(c, "F2F5F9")
    for row in t.rows:
        for k, w in enumerate(widths): row.cells[k].width = w
    d.add_paragraph()
    return t


# ═══════════════════════════ TRAILER ═══════════════════════════

tl, tr_total = timeline(TR)
TRAILER = {
    "title": ("Mở đầu — Logo", "Thẻ tĩnh, logo fade-in", "Logo NihongoLife, tagline 「ひばり町で、日本語と暮らそう」", "—", "Pad nhạc mở (D yo pentatonic), chưa có nhịp"),
    "city": ("Toàn cảnh Hibari", "Flycam/crane cao, trượt ngang và hạ dần", "Toàn cảnh mái nhà, đường phố, biển hiệu tiếng Nhật", "ひばり町へ、ようこそ。\n(Chào mừng đến thị trấn Hibari.)", "Nhạc vào nhịp, kick nhẹ"),
    "walk": ("Dạo phố", "Tracking sau lưng nhân vật, ngang vai", "Nhân vật đi bộ qua biển chỉ đường song ngữ", "歩いて、見つけて、話そう。\n(Đi, khám phá, trò chuyện.)", "Tiếng bước chân hòa nhạc"),
    "dialogue": ("Hội thoại N5", "Qua vai (OTS) về phía NPC trước Game Center", "Hộp hội thoại 3 lựa chọn, có furigana + nghĩa tiếng Việt", "場面に合った言葉を、選んで答える。\n(Chọn câu phù hợp với tình huống.)", "Nhạc giữ nhịp, giọng đọc nổi"),
    "konbini": ("Mua sắm ở konbini", "Trung cảnh trong cửa hàng → cửa sổ mua hàng", "Kệ hàng 3D, cửa sổ chọn onigiri có giá yên", "コンビニで、買い物。\n(Mua sắm ở cửa hàng tiện lợi.)", "Pluck melody vào"),
    "job": ("Làm thêm", "OTS sau nhân vật, đẩy máy (push-in) chậm tới quầy", "Ca làm thêm: khách hỏi bằng tiếng Nhật, thẻ câu hỏi 4 lựa chọn, chọn đúng", "アルバイトで、働こう。\n(Đi làm thêm thôi.)", "Nhạc đầy đủ"),
    "journal": ("Sổ nhiệm vụ", "Toàn cảnh phố, cửa sổ sổ nhiệm vụ phía trước", "Tab Làm thêm: mục tiêu, lương ¥600, XP, trạng thái ca", "タスクをこなして、レベルアップ。\n(Làm nhiệm vụ, lên cấp.)", "—"),
    "station": ("Ga Hibari", "Dolly chậm dọc sân ga", "Sân ga, bảng tên ga, đoàn tàu", "駅で、切符を買って。\n(Mua vé ở nhà ga.)", "Chuông nhà ga hòa vào nhạc"),
    "train": ("Lên tàu", "Góc trong toa / cửa sổ", "Hành trình tới đảo Midori", "電車で、みどり島へ。\n(Đi tàu tới đảo Midori.)", "Chuyển đoạn (build-up)"),
    "island": ("Đảo Midori", "Crane từ trên cao hạ xuống toàn đảo", "Toàn đảo: ruộng, chuồng, ga, biển; bảng chào mừng", "海に浮かぶ、みどり島。\n(Đảo Midori giữa biển.)", "Nhạc mở rộng, pad sáng"),
    "farm": ("Làm nông", "Góc 3/4 trung cảnh, push-in nhẹ", "Xới (くわ) → gieo → tưới (じょうろ); thanh tiến trình theo dụng cụ", "耕して、植えて、水をあげる。\n(Xới, gieo, tưới.)", "Nhịp đều theo thao tác"),
    "grow": ("Cây lớn", "Cận cảnh thấp, máy lùi chậm; time-lapse", "Cà rốt qua 3 giai đoạn (đồng hồ nông trại tua nhanh)", "育てて、収穫。\n(Chăm sóc và thu hoạch.)", "Chime nhẹ mỗi giai đoạn"),
    "animals": ("Vật nuôi", "Toàn cảnh chuồng → trung cảnh, thẻ con vật", "Bò, alpaca, lừa, ngựa; thẻ con vật với từ vựng 「うし」", "動物たちと、仲良く。\n(Thân thiết với các con vật.)", "—"),
    "sell": ("Bán nông sản", "Trung cảnh quầy, cửa sổ cửa hàng", "Tab Bán: bán cà rốt, ví chung tăng tiền", "売って、買って、暮らす。\n(Mua bán và sinh sống.)", "—"),
    "classroom": ("Lớp học", "Toàn cảnh lớp học từ cuối phòng", "Hibari Nihongo Gakuin: lớp học, phòng thi thử", "ひばり日本語学院で、試験に挑戦。\n(Thử sức kỳ thi ở học viện Hibari.)", "Nhạc hạ nhẹ"),
    "kana": ("Kana Match", "Toàn màn hình mini-game", "Lật thẻ, ghép cặp kana–romaji", "遊びながら、かなを覚えよう。\n(Vừa chơi vừa nhớ kana.)", "Nhịp nhanh, chuẩn bị kết"),
    "end": ("Kết — End card", "Thẻ tĩnh, chữ fade-in theo lớp", "Logo, tagline, danh sách tính năng, Capstone 2026 · VTC Academy", "ひばり町で、日本語と暮らそう。\n(Sống cùng tiếng Nhật ở Hibari.)", "Swell kết + chuông gió, fade-out 3 s"),
}
ACTS = {"title": "Hồi 1 — Đến Hibari", "job": "Hồi 2 — Làm việc & lên đường", "island": "Hồi 3 — Cuộc sống trên đảo", "classroom": "Hồi 4 — Học & kết"}

d = new_doc()
title_block(d, "NIHONGOLIFE · ĐỒ ÁN CAPSTONE 2026", "KỊCH BẢN TRAILER (BẢN 2)",
            f"Bản dựng khớp video NihongoLife_Trailer_v2.mp4 · thời lượng {tc(tr_total)} · 1920×1080, 30 fps")
d.add_heading("1. Thông tin chung", 1)
table(d, ["Hạng mục", "Nội dung"], [
    ["Mục đích", "Giới thiệu NihongoLife trong khoảng một phút rưỡi: một người học tiếng Nhật đến thị trấn Hibari, giao tiếp, đi làm thêm, đi tàu ra đảo Midori làm nông, rồi quay lại lớp học và Game Center."],
    ["Thời lượng", f"{tc(tr_total)} ({tr_total} khung hình, 30 fps), {len(tl)} cảnh"],
    ["Nguồn hình", "Ghi trực tiếp trong Unity 6 (URP) bằng TrailerReel: mỗi cảnh do mã điều khiển camera và gọi đúng logic gameplay (mua hàng, ca làm, ruộng, cửa hàng); Time.captureFramerate = 30 nên không rớt khung."],
    ["Âm nhạc", "Nhạc nền tự tổng hợp bằng mã (thang D yo pentatonic, 92 BPM): pad, bass, kick, hi-hat và giai điệu pluck; không dùng nhạc có bản quyền."],
    ["Lời dẫn", "Tiếng Nhật, giọng đọc AI (Microsoft Edge TTS, ja-JP-NanamiNeural). Phụ đề tiếng Nhật + tiếng Việt hiển thị ngay trong game. Nhạc tự hạ khi có lời (sidechain)."],
    ["Đối tượng", "Hội đồng chấm Capstone, người học tiếng Nhật trình độ N5, người chơi thích game mô phỏng đời sống."],
], [Cm(3.5), Cm(13)])

d.add_heading("2. Ý tưởng và mạch kể chuyện", 1)
para(d, "Trailer kể một “ngày sống ở Nhật” theo bốn hồi, mỗi hồi mở ra một vòng lặp gameplay đã có trong bản hiện tại. Mỗi cảnh chỉ mang một thông điệp, viết thành một câu tiếng Nhật ngắn, trình độ N5. Như vậy trailer cũng là một bài nghe nhỏ cho người học.", align="j")
bullets(d, [
    "**Hồi 1 — Đến Hibari:** thị trấn, đi bộ, hội thoại có lựa chọn, mua đồ ở konbini.",
    "**Hồi 2 — Làm việc và lên đường:** ca làm thêm, sổ nhiệm vụ (lương, cấp độ), ga tàu, lên tàu.",
    "**Hồi 3 — Cuộc sống trên đảo:** toàn cảnh đảo Midori, xới–gieo–tưới bằng dụng cụ, cây lớn, vật nuôi, mua bán.",
    "**Hồi 4 — Học và kết:** lớp học/phòng thi thử, mini-game Kana Match, end card.",
])
para(d, "Nhịp dựng: các cảnh dài 4–6,5 giây và chuyển bằng cắt thẳng theo phách nhạc. Các cảnh có thao tác (làm thêm, làm nông) dài hơn để người xem kịp thấy thanh tiến trình và kết quả.", align="j")

landscape(d)
d.add_heading("3. Shot list (timecode theo bản dựng)", 1)
rows = []
for idx, (sid, a, b) in enumerate(tl, 1):
    name, cam, content, line, sound = TRAILER.get(sid, (sid, "", "", "", ""))
    act = ACTS.get(sid)
    rows.append([str(idx), f"{tc(a)}–{tc(b)}\n({(b - a) / 30:.1f}s)", thumb(TR, a, b, "tr_" + sid),
                 (f"**{act}**\n" if act else "") + f"**{name}**", cam, content, line, sound])
table(d, ["#", "Timecode", "Khung hình", "Cảnh", "Góc máy / chuyển động", "Nội dung & hiệu ứng", "Lời dẫn (JA) / nghĩa", "Âm thanh"],
      rows, [Cm(0.8), Cm(2.0), Cm(4.0), Cm(3.2), Cm(3.8), Cm(5.0), Cm(4.6), Cm(3.1)], size=9, images={2})

portrait(d)
d.add_heading("4. Hiệu ứng hình và chữ", 1)
bullets(d, [
    "**Phụ đề trong game:** dòng tiếng Nhật (to) + dòng tiếng Việt (nhỏ) ở cạnh dưới; fade-in 0,45 s, giữ hết cảnh, fade-out trước khi cắt.",
    "**Camera:** đường bay nội suy mượt (ease in/out) giữa hai điểm đặt camera và hai điểm nhìn; không rung tay, không zoom số.",
    "**Giao diện thật:** mọi cửa sổ (hội thoại, cửa hàng, sổ nhiệm vụ, thẻ câu hỏi, thẻ ruộng) là UI đang chạy trong game, không ghép ảnh.",
    "**Time-lapse cây lớn:** đồng hồ nông trại được tua nhanh bằng mã. Ngoài trailer, cây lớn theo thời gian thực (20–50 giây mỗi giai đoạn tùy loại cây).",
    "**Thẻ mở/kết:** logo trên nền xanh đậm, chữ hiện dần theo lớp; thẻ kết liệt kê bảy nhóm tính năng.",
])
d.add_heading("5. Thiết kế âm thanh", 1)
table(d, ["Đoạn", "Âm nhạc", "Lời / hiệu ứng"], [
    ["Mở đầu", "Pad + bass, chưa có trống (khoảng 4 giây đầu)", "Không lời dẫn trên logo"],
    ["Thân trailer", "Kick ở phách 1 và 3, hi-hat móc, giai điệu pluck pentatonic; thay đổi hợp âm theo ô nhịp", "Mỗi cảnh một câu tiếng Nhật, bắt đầu 0,4 s sau điểm cắt; nhạc hạ khoảng 6 dB khi có lời"],
    ["Kết", "Swell pad và chuông gió ở 10 giây cuối, fade-out 3 giây", "Câu tagline cuối"],
], [Cm(3), Cm(7), Cm(6.5)])
d.add_heading("6. Quy trình dựng (tái lập được)", 1)
bullets(d, [
    "Ghi hình: chạy PlayMode test Explicit **TrailerRecordingTests.Trailer_RecordFrames** (biến môi trường NL_FRAME_DIR chọn thư mục). Kết quả là ảnh JPG 1920×1080 từng khung và timeline.tsv ghi khung bắt đầu của mỗi cảnh.",
    "Nhạc: **trailer_music.py <wav> <độ dài>**; lời dẫn: edge-tts từng câu trong tts_trailer/lines.tsv.",
    "Trộn và xuất: **mix_video.py** đặt từng câu theo timeline, nén nhạc theo giọng (sidechain), xuất H.264 CRF 19 + AAC 192 kbps.",
    "Kiểm tra: trích khung giữa mỗi cảnh thành contact sheet và xem lại toàn bộ video trước khi nộp.",
])
d.add_heading("7. Ghi chú trung thực", 1)
bullets(d, [
    "Mọi cảnh lấy từ bản Unity hiện tại; không có cảnh dựng giả hay tính năng chưa làm.",
    "Lời dẫn là giọng đọc AI (TTS), không phải diễn viên lồng tiếng.",
    "Trong cảnh “Cây lớn”, thời gian được tua nhanh có chủ đích, như đã ghi ở mục 4.",
])
d.save(OUT / "NihongoLife_Kich_Ban_Trailer_v2.docx")

# ═══════════════════════════ DEMO ═══════════════════════════

dl, de_total = timeline(DE)
lines = dict(l.split("\t", 1) for l in (HERE / "tts_demo" / "lines.tsv").read_text(encoding="utf-8").split("\n") if "\t" in l)
DEMO = {
    "intro": ("Mở đầu", "— (thẻ tiêu đề, camera xoay quanh nhân vật)", "Giới thiệu sản phẩm, nguồn ghi hình"),
    "hud": ("1 · Giao diện", "—", "HUD thu gọn: vòng trạng thái, 5 chỉ số, tiền; thông báo tự ẩn"),
    "move": ("1 · Giao diện", "W A S D đi bộ · chuột xoay camera · giữ Ctrl hiện con trỏ", "Điều khiển mouse-look, Ctrl giải phóng con trỏ"),
    "settings": ("1 · Giao diện", "O mở Cài đặt · Esc đóng", "Âm lượng, ngôn ngữ, độ nhạy chuột, đảo trục, đổi phím"),
    "profile": ("1 · Giao diện", "Tab mở hồ sơ · Esc đóng", "Cấp độ, XP, điểm kiến thức tách riêng"),
    "journal": ("1 · Giao diện", "N mở sổ nhiệm vụ · bấm “Nhận ca làm”", "Quest/Job: mục tiêu, lương, nhận việc"),
    "stock": ("2 · Konbini", "F ở thùng kho → F ở 3 kệ", "Thao tác có thanh tiến trình, mang thùng hàng"),
    "customer": ("2 · Konbini", "F ở khách → chọn sai → chọn đúng", "Câu hỏi tiếng Nhật; sai không tính điểm"),
    "customer2": ("2 · Konbini", "F ở khách thứ hai → chọn đúng", "Mục tiêu hỗ trợ khách 2/2"),
    "register": ("2 · Konbini", "F ở quầy → quét → chọn tổng tiền", "Tính tiền bằng tiếng Nhật"),
    "paid": ("2 · Konbini", "F nói chuyện với Ito", "Lương ¥600, +40 XP, +6 kiến thức, trả một lần"),
    "spend": ("2 · Konbini", "Chọn onigiri → Thêm vào giỏ → Thanh toán", "Ví chung, nhiệm vụ hằng ngày hoàn thành"),
    "station": ("3 · Ga & tàu", "F máy bán vé → chọn Midori (¥450) → F cổng soát vé", "Mua vé, soát vé"),
    "train": ("3 · Ga & tàu", "F lên tàu khi tàu mở cửa", "Chuyển vùng bằng tàu"),
    "island": ("4 · Đảo Midori", "—", "Toàn cảnh đảo, ví chung"),
    "hana": ("4 · Đảo Midori", "F nói chuyện với Hana → “Nhận ca làm”", "Job nông trại, được phát 2 gói hạt"),
    "hand": ("4 · Đảo Midori", "F ở luống 1 (chưa có cuốc)", "Không có dụng cụ thì không xới được; thẻ ô ruộng mở cửa hàng"),
    "hoe": ("4 · Đảo Midori", "Cầm cuốc → F ở luống 2 → “Xới đất”", "Xới 1,6 s; cuốc mòn 1/40 mỗi lần, hết bền thì hỏng"),
    "plant": ("4 · Đảo Midori", "“Gieo cà rốt” → “Tưới nước” (2 luống)", "Gieo, tưới cần bình tưới"),
    "grow": ("4 · Đảo Midori", "Tưới lại khi cây cần nước", "Cây lớn theo thời gian thực (video tua nhanh)"),
    "harvest": ("4 · Đảo Midori", "“Thu hoạch” → thẻ vật nuôi → “Cho ăn”", "Nông sản dùng làm thức ăn"),
    "farmpaid": ("4 · Đảo Midori", "F báo cáo với Hana", "Tiền công ¥500"),
    "sell": ("4 · Đảo Midori", "F ở quầy → tab Bán → Xác nhận", "Bán nông sản vào ví chung"),
    "english": ("4 · Đảo Midori", "Chuyển ngôn ngữ mục tiêu JA ↔ EN", "Nội dung song ngữ JA/EN"),
    "save": ("4 · Đảo Midori", "Tab mở hồ sơ", "Lưu tiến trình (tiền, XP, nhiệm vụ, túi đồ, ruộng)"),
    "return": ("5 · Trở về", "F máy bán vé đảo → F cửa tàu", "Vé về Hibari, màn hình di chuyển"),
    "back": ("5 · Trở về", "—", "Về ga Hibari, người chơi điều khiển lại được"),
    "outro": ("Kết", "— (thẻ kết)", "Các phần khác: lớp học, thi thử, Kana Match"),
}

d = new_doc()
title_block(d, "NIHONGOLIFE · ĐỒ ÁN CAPSTONE 2026", "KỊCH BẢN VIDEO DEMO GAMEPLAY (BẢN 2)",
            f"Bản dựng khớp video NihongoLife_Gameplay_Demo.mp4 · thời lượng {tc(de_total)} · 1920×1080, 30 fps")
d.add_heading("1. Mục tiêu", 1)
para(d, "Video demo chứng minh trước hội đồng rằng vòng lặp cuộc sống – việc làm – kinh tế của NihongoLife hoạt động trọn vẹn trong một lượt chơi liên tục: điều khiển và giao diện, một ca làm thêm có câu hỏi tiếng Nhật, lĩnh lương rồi tiêu tiền, đi tàu ra đảo Midori, làm ca nông trại, mua bán, đổi ngôn ngữ mục tiêu và trở về. Mỗi bước là logic gameplay thật, trùng với kịch bản kiểm thử tự động LifeLoopPlayModeTests.", align="j")
table(d, ["Hạng mục", "Nội dung"], [
    ["Thời lượng", f"{tc(de_total)} · {len(dl)} đoạn thuyết minh, 5 phần + mở/kết"],
    ["Góc nhìn", "Camera gameplay góc thứ ba, giữ nguyên HUD. Con trỏ (vòng vàng) và phím vừa bấm (ô phím bên trái) được hiển thị để người xem theo dõi thao tác."],
    ["Thuyết minh", "Tiếng Việt, giọng đọc AI (Edge TTS, vi-VN-HoaiMyNeural), kèm phụ đề tiếng Việt."],
    ["Âm nhạc", "Nền nhạc nhẹ tự tổng hợp (cùng chủ đề với trailer, bỏ trống), tự hạ khi có lời."],
    ["Trạng thái ban đầu", "Người chơi mới ở Hibari với ¥300, chưa có nhiệm vụ và dụng cụ; đảo Midori ở trạng thái mới."],
], [Cm(3.5), Cm(13)])

landscape(d)
d.add_heading("2. Trình tự thao tác và lời thuyết minh", 1)
rows = []
for idx, (sid, a, b) in enumerate(dl, 1):
    part, keys, feat = DEMO.get(sid, (sid, "", ""))
    rows.append([str(idx), f"{tc(a)}\n({(b - a) / 30:.0f}s)", thumb(DE, a, b, "de_" + sid, 0.6), f"**{part}**", keys, feat, lines.get(sid, "")])
table(d, ["#", "Bắt đầu", "Khung hình", "Phần", "Thao tác (phím / chuột)", "Tính năng chứng minh", "Lời thuyết minh"],
      rows, [Cm(0.8), Cm(1.7), Cm(3.6), Cm(2.3), Cm(4.6), Cm(4.4), Cm(9.0)], size=9, images={2})

portrait(d)
d.add_heading("3. Hướng dẫn ghi hình", 1)
d.add_heading("3.1. Bản tự động (đã dùng cho video này)", 2)
bullets(d, [
    "Đóng Unity Editor, rồi chạy PlayMode test Explicit **DemoRecordingTests.Demo_RecordFrames** ở chế độ batchmode. Biến NL_DEMO_DIR chọn thư mục khung hình; NL_DEMO_VO trỏ tới tts_demo/durations.tsv để mỗi đoạn kéo dài ít nhất bằng câu thuyết minh của nó.",
    "Test vừa ghi hình vừa kiểm tra: ca konbini phải hoàn thành, ca nông trại phải đủ mục tiêu và được trả lương, vé tàu phải hợp lệ. Nếu một bước hỏng, test báo lỗi thay vì xuất video sai.",
    "Dựng: **mix_video.py <frames> tts_demo bed_demo.wav NihongoLife_Gameplay_Demo.mp4 --subs --mvol 0.35**.",
])
d.add_heading("3.2. Bản trình diễn trực tiếp khi bảo vệ", 2)
bullets(d, [
    "**Phần mềm:** OBS Studio, canvas 1920×1080, 30 fps, encoder x264/NVENC CQ 20; nguồn Game Capture (hoặc Window Capture cửa sổ Unity ở chế độ Play, Game view 1920×1080, tắt Stats/Gizmos).",
    "**Âm thanh:** micro thu giọng thuyết minh khoảng −12 dB; âm game khoảng −20 dB; bật Noise Suppression.",
    "**Chuẩn bị save:** tạo người chơi mới (khách) hoặc dùng save sạch: ¥300, chưa có nhiệm vụ. Kiểm tra trước ca konbini đang ở trạng thái “Có thể nhận”.",
    "**Cài đặt:** chế độ điều khiển Mouse-look, độ nhạy 1.0×; âm lượng nhạc 40% để giọng nói rõ.",
    "**Đi theo đúng bảng ở mục 2.** Mỗi lần mở cửa sổ, dừng khoảng 1 giây trước khi bấm để người xem kịp đọc.",
    "**Chờ tàu:** tàu chạy theo lịch. Nếu tàu chưa vào ga, nói về sổ nhiệm vụ hoặc ví tiền trong lúc chờ, không cắt cảnh.",
    "**Cây lớn:** ở bản trực tiếp, gieo hạt trước khi trình bày phần khác rồi quay lại thu hoạch (20–50 giây mỗi giai đoạn), hoặc nói rõ là đang dùng bản ghi có tua nhanh.",
])
d.add_heading("3.3. Phương án dự phòng", 2)
bullets(d, [
    "Nếu máy trình chiếu yếu hoặc Unity lỗi, phát video NihongoLife_Gameplay_Demo.mp4 có sẵn rồi trả lời câu hỏi trên Editor sau.",
    "Nếu mất mạng: toàn bộ demo chạy offline; tiến trình lưu cục bộ.",
    "Giữ sẵn các ảnh chụp kiểm thử trong Bao_Cao/lifeloop-regression và island-regression làm bằng chứng phụ.",
])
d.add_heading("4. Ghi chú trung thực", 1)
bullets(d, [
    "Ở bước “chưa có cuốc”, cuốc trong bộ khởi đầu được cất đi để minh họa trường hợp thiếu dụng cụ, rồi trả lại ngay ở bước sau.",
    "Ở bước “cây lớn”, đồng hồ nông trại được tua nhanh; lời thuyết minh nói rõ điều này.",
    "Lời thuyết minh và lời dẫn là giọng đọc AI (TTS).",
    "Video được ghi trong Unity Editor (PlayMode); bản build WebGL chưa được kiểm thử.",
])
d.save(OUT / "NihongoLife_Kich_Ban_Demo_v2.docx")
print("scripts written")
