import os
import sys

# Let's inspect installed modules or write an openpyxl script that generates an ultra-rich XLSX file.

script_content = '''
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

wb = openpyxl.Workbook()

# Sheet 1: Timeline & Progress Tracking
ws1 = wb.active
ws1.title = "Timeline & Progress"
ws1.views.sheetView[0].showGridLines = True

# Title Header Banner
ws1.merge_cells("A1:M2")
title_cell = ws1["A1"]
title_cell.value = "KẾ HOẠCH & TIẾN ĐỘ DỰ ÁN NIHONGOLIFE (SOLO DEVELOPER: QUÁCH THÀNH LONG)"
title_cell.font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Slate / Navy
title_cell.alignment = Alignment(horizontal="center", vertical="center")

# Subtitle Info Bar
ws1.merge_cells("A3:M3")
sub_cell = ws1["A3"]
sub_cell.value = "Dự án: NihongoLife (E-learning JLPT + IELTS 3D Metaverse) | Tác giả: Quách Thành Long | Thời gian: 01/06/2026 - 15/10/2026"
sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="475569")
sub_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")

# Table Headers
headers = [
    "STT", "Mã Đấu Công Việc", "Giai Đoạn (Phase)", "Đầu Công Việc Chi Tiết", 
    "Hạng Mục / Chức Năng", "Mức Độ Ưu Tiên", "Ngày Bắt Đầu", "Deadline (Hạn Chót)", 
    "Thời Gian (Ngày)", "Trạng Thái", "Độ Hoàn Thành (%)", "Dấu Tích", "Ghi Chú & Kết Quả"
]

header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

ws1.row_dimensions[5].height = 28

for col_idx, h in enumerate(headers, 1):
    cell = ws1.cell(row=5, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

# Data Definition
tasks = [
    # Phase 1: Planning & Setup (01/06 - 15/06)
    (1, "TASK-001", "Phase 1: Ý Tưởng & Khảo Sát", "Nghiên cứu thị trường & Đề xuất dự án", "Khảo sát nhu cầu học JLPT/IELTS qua 3D Game", "Cao", "2026-06-01", "2026-06-05", 5, "Đã hoàn thành", 1.0, "☑ Đạt", "Đã phê duyệt đề tài Capstone Project 3"),
    (2, "TASK-002", "Phase 1: Ý Tưởng & Khảo Sát", "Viết SRS & Game Design Document (GDD)", "Xác định Game Mechanics, Core Loop, SRS UI/UX", "Rất cao", "2026-06-06", "2026-06-12", 7, "Đã hoàn thành", 1.0, "☑ Đạt", "Hoàn thành bộ tài liệu GDD & SRS 100%"),
    (3, "TASK-003", "Phase 1: Ý Tưởng & Khảo Sát", "Thiết lập Repository & Unity Project", "Cấu hình Git, Unity 6000.3.12f1, URP Pipeline", "Trung bình", "2026-06-13", "2026-06-15", 3, "Đã hoàn thành", 1.0, "☑ Đạt", "Khởi tạo Git repo và base project cấu trúc chuẩn"),

    # Phase 2: Core Architecture & Database (16/06 - 05/07)
    (4, "TASK-004", "Phase 2: Kiến Trúc & Supabase", "Tích hợp Cloud Supabase Backend", "Thiết lập Auth, REST Client, Leaderboard API", "Rất cao", "2026-06-16", "2026-06-22", 7, "Đã hoàn thành", 1.0, "☑ Đạt", "SupabaseClient.cs chạy ổn định cloud persistence"),
    (5, "TASK-005", "Phase 2: Kiến Trúc & Supabase", "Xây dựng Save/Load Progress System", "Local JSON fallback + Cloud sync song song", "Cao", "2026-06-23", "2026-06-28", 6, "Đã hoàn thành", 1.0, "☑ Đạt", "IProgressRepository hoạt động không lỗi"),
    (6, "TASK-006", "Phase 2: Kiến Trúc & Supabase", "Quản lý dữ liệu Scenario & Quests Data", "Thiết kế ScriptableObjects, JSON quest flow", "Cao", "2026-06-29", "2026-07-05", 7, "Đã hoàn thành", 1.0, "☑ Đạt", "Nội dung quest & câu thoại data-driven"),

    # Phase 3: World & Environment Design (06/07 - 31/07)
    (7, "TASK-007", "Phase 3: Thế Giới 3D & Scene", "Thiết kế Map Station District (Ga Tàu)", "Xây dựng 20_StationDistrict, NavMesh, Lighting", "Cao", "2026-07-06", "2026-07-15", 10, "Đã hoàn thành", 1.0, "☑ Đạt", "Cảnh quan đẹp mắt, tối ưu URP batching"),
    (8, "TASK-008", "Phase 3: Thế Giới 3D & Scene", "Thiết kế Map Sushi Restaurant (Nhà Hàng)", "Xây dựng 30_SushiRestaurant, NPC Kaiwa order", "Trung bình", "2026-07-16", "2026-07-23", 8, "Đã hoàn thành", 1.0, "☑ Đạt", "NPC hội thoại JLPT N5/N4 nhà hàng"),
    (9, "TASK-009", "Phase 3: Thế Giới 3D & Scene", "Tối ưu hóa Sandbox & Bootstrap Flow", "Streamline runtime load thẳng 90_TestSandbox", "Cao", "2026-07-24", "2026-07-31", 8, "Đã hoàn thành", 1.0, "☑ Đạt", "Bypass MainMenu boost tốc độ test 5x"),

    # Phase 4: E-Learning Engine (JLPT + IELTS) (01/08 - 25/08)
    (10, "TASK-010", "Phase 4: Hệ Thống E-Learning", "Xây dựng Assessment Models & Question Bank", "Question Blueprint, seed ngẫu nhiên, auto grading", "Rất cao", "2026-08-01", "2026-08-08", 8, "Đã hoàn thành", 1.0, "☑ Đạt", "AssessmentEngine.cs & AssessmentModels.cs complete"),
    (11, "TASK-011", "Phase 4: Hệ Thống E-Learning", "Phát triển JLPT Practice & Review Flow", "Chấm điểm JLPT N5-N1, lưu câu sai làm lại", "Rất cao", "2026-08-09", "2026-08-16", 8, "Đã hoàn thành", 1.0, "☑ Đạt", "ExamCenterPopup & ExamManager hoàn thiện"),
    (12, "TASK-012", "Phase 4: Hệ Thống E-Learning", "Phát triển IELTS 4 Skills Engine", "Mod modules Listening/Speaking/Reading/Writing", "Cao", "2026-08-17", "2026-08-25", 9, "Đã hoàn thành", 1.0, "☑ Đạt", "Tính điểm Band Score chuẩn 9.0 scale"),

    # Phase 5: Metaverse & Advanced Integration (26/08 - 15/09)
    (13, "TASK-013", "Phase 5: Metaverse & Advanced", "Tích hợp SDK Video Call 3D (Agora Native)", "EduMeetingManager.cs, map stream lên TV screen 3D", "Cao", "2026-08-26", "2026-09-03", 9, "Đã hoàn thành", 1.0, "☑ Đạt", "Video call trong không gian VR/3D thành công"),
    (14, "TASK-014", "Phase 5: Metaverse & Advanced", "Phát triển Learning Center (Trường Học)", "Build 50_LearningCenter với asset StylooClassroom", "Rất cao", "2026-09-04", "2026-09-10", 7, "Đã hoàn thành", 1.0, "☑ Đạt", "Scene 50_LearningCenter đầy đủ phòng học"),
    (15, "TASK-015", "Phase 5: Metaverse & Advanced", "Phát triển Shopping District (Trung Tâm TM)", "Build 40_ShoppingDistrict mua sắm vật phẩm", "Trung bình", "2026-09-11", "2026-09-15", 5, "Đã hoàn thành", 1.0, "☑ Đạt", "Scene 40_ShoppingDistrict hoàn thành"),

    # Phase 6: Testing, Refactoring & Fixes (16/09 - 30/09)
    (16, "TASK-016", "Phase 6: Kiểm Thử & Sửa Lỗi", "Chạy PlayMode & EditMode Tests", "Automated suite test collision, NPC, zone portals", "Rất cao", "2026-09-16", "2026-09-22", 7, "Đã hoàn thành", 1.0, "☑ Đạt", "Pass 100% test scenario regression"),
    (17, "TASK-017", "Phase 6: Kiểm Thử & Sửa Lỗi", "Fix Unity Licensing & Batch Build Issue", "Khắc phục kết nối Licensing Client trên Windows", "Cao", "2026-09-23", "2026-09-26", 4, "Đã hoàn thành", 1.0, "☑ Đạt", "Batch mode build tự động hóa trơn tru"),
    (18, "TASK-018", "Phase 6: Kiểm Thử & Sửa Lỗi", "Sửa lỗi UI Map & Navigation Pathing", "WorldMapUI click tab & Player Click-to-Move NavMesh", "Trung bình", "2026-09-27", "2026-09-30", 4, "Đã hoàn thành", 1.0, "☑ Đạt", "Di chuyển mượt mà không dính tường"),

    # Phase 7: Document, Media & Presentation (01/10 - 15/10)
    (19, "TASK-019", "Phase 7: Báo Cáo & Bảo Vệ", "Quay Video Trailer & Gameplay Showcase", "Record NihongoLife_Trailer.mp4 4K chất lượng cao", "Cao", "2026-10-01", "2026-10-03", 3, "Đã hoàn thành", 1.0, "☑ Đạt", "Trailer ấn tượng dài 3 phút trong Bao_Cao"),
    (20, "TASK-020", "Phase 7: Báo Cáo & Bảo Vệ", "Viết Báo Cáo Capstone Project 3 (Word/PDF)", "NihongoLife_Bao_Cao_Capstone.docx (>100 trang)", "Rất cao", "2026-10-04", "2026-10-06", 3, "Đã hoàn thành", 1.0, "☑ Đạt", "Báo cáo chi tiết cấu trúc kiến trúc & kết quả"),
    (21, "TASK-021", "Phase 7: Báo Cáo & Bảo Vệ", "Thiết kế Slide Thuyết Trình (PowerPoint)", "NihongoLife_Thuyet_Trinh_moi.pptx thiết kế đẹp", "Cao", "2026-10-06", "2026-10-07", 2, "Đã hoàn thành", 1.0, "☑ Đạt", "Slide 35 trang hiện đại chuyên nghiệp"),
    (22, "TASK-022", "Phase 7: Báo Cáo & Bảo Vệ", "Tạo Excel Timeline & Tiến Độ (Solo Dev)", "Excel Timeline chi tiết phân bổ hạn chót & đánh giá", "Cao", "2026-10-08", "2026-10-08", 1, "Đang thực hiện", 0.95, "☐ Đang làm", "Tệp Excel xịn đẹp siêu chỉn chu"),
    (23, "TASK-023", "Phase 7: Báo Cáo & Bảo Vệ", "Tổng Duyệt & Đóng Gói Sản Phẩm (Release)", "Kiểm tra WebGL build, kiểm tra tổng thể đồ án", "Rất cao", "2026-10-09", "2026-10-15", 7, "Chưa bắt đầu", 0.0, "☐ Chưa đạt", "Chuẩn bị sẵn sàng bảo vệ đồ án Capstone")
]

start_row = 6
for r_idx, task in enumerate(tasks, start=start_row):
    ws1.row_dimensions[r_idx].height = 22
    for c_idx, val in enumerate(task, start=1):
        cell = ws1.cell(row=r_idx, column=c_idx, value=val)
        cell.font = Font(name="Segoe UI", size=10)
        cell.border = thin_border
        
        # Alignment & Formatting
        if c_idx in [1, 2, 7, 8, 9, 10, 12]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c_idx in [11]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '0.0%'
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
        # Priority Colors
        if c_idx == 6:
            if val == "Rất cao":
                cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Soft red
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
            elif val == "Cao":
                cell.fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid") # Soft orange
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="9A3412")
            elif val == "Trung bình":
                cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Soft yellow
                cell.font = Font(name="Segoe UI", size=10, color="92400E")

        # Status Colors
        if c_idx == 10:
            if val == "Đã hoàn thành":
                cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Soft green
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
            elif val == "Đang thực hiện":
                cell.fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid") # Soft blue
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="075985")
            elif val == "Chưa bắt đầu":
                cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
                cell.font = Font(name="Segoe UI", size=10, italic=True, color="6B7280")
                
        # Checkmark Icon styling
        if c_idx == 12:
            if "☑" in str(val):
                cell.font = Font(name="Segoe UI", size=11, bold=True, color="16A34A")
            else:
                cell.font = Font(name="Segoe UI", size=11, color="EA580C")

# Summary Section Below Table
sum_row = start_row + len(tasks) + 1
ws1.merge_cells(f"A{sum_row}:C{sum_row}")
ws1.cell(row=sum_row, column=1, value="TỔNG CỘNG HẠNG MỤC").font = Font(name="Segoe UI", bold=True, color="0F172A")
ws1.cell(row=sum_row, column=1).alignment = Alignment(horizontal="center", vertical="center")
ws1.cell(row=sum_row, column=1).fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

total_tasks_cell = ws1.cell(row=sum_row, column=4, value=f"=COUNTA(A6:A{sum_row-2})")
total_tasks_cell.font = Font(name="Segoe UI", bold=True)
total_tasks_cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

ws1.cell(row=sum_row, column=8, value="TỔNG NGÀY:").font = Font(name="Segoe UI", bold=True)
ws1.cell(row=sum_row, column=8).alignment = Alignment(horizontal="right", vertical="center")
ws1.cell(row=sum_row, column=8).fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

total_days_cell = ws1.cell(row=sum_row, column=9, value=f"=SUM(I6:I{sum_row-2})")
total_days_cell.font = Font(name="Segoe UI", bold=True, color="1E40AF")
total_days_cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
total_days_cell.alignment = Alignment(horizontal="center", vertical="center")

ws1.cell(row=sum_row, column=10, value="TIẾN ĐỘ CHUNG:").font = Font(name="Segoe UI", bold=True)
ws1.cell(row=sum_row, column=10).alignment = Alignment(horizontal="right", vertical="center")
ws1.cell(row=sum_row, column=10).fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

avg_progress_cell = ws1.cell(row=sum_row, column=11, value=f"=AVERAGE(K6:K{sum_row-2})")
avg_progress_cell.font = Font(name="Segoe UI", bold=True, color="15803D")
avg_progress_cell.number_format = '0.0%'
avg_progress_cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
avg_progress_cell.alignment = Alignment(horizontal="right", vertical="center")

# Border for summary row
for c in range(1, 14):
    ws1.cell(row=sum_row, column=c).border = thin_border

# Sheet 2: Dashboard Overview & Thống Kê
ws2 = wb.create_sheet(title="Dashboard KPI")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:G2")
d_title = ws2["A1"]
d_title.value = "THỐNG KÊ & TỔNG QUAN TIẾN ĐỘ SOLO DEVELOPER"
d_title.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
d_title.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
d_title.alignment = Alignment(horizontal="center", vertical="center")

# KPI Summary Cards
kpis = [
    ("TỔNG ĐẦU CÔNG VIỆC", len(tasks), "1E40AF", "EFF6FF"),
    ("ĐÃ HOÀN THÀNH", len([t for t in tasks if t[9] == "Đã hoàn thành"]), "166534", "F0FDF4"),
    ("ĐANG THỰC HIỆN", len([t for t in tasks if t[9] == "Đang thực hiện"]), "075985", "F0F9FF"),
    ("CHƯA BẮT ĐẦU", len([t for t in tasks if t[9] == "Chưa bắt đầu"]), "374151", "F9FAFB"),
    ("TỔNG THỜI GIAN", f"{sum([t[8] for t in tasks])} Ngày", "6B21A8", "F3E8FF"),
]

for idx, (kpi_title, kpi_val, color_hex, bg_hex) in enumerate(kpis, start=1):
    col_letter = get_column_letter(idx * 2 - 1)
    next_col = get_column_letter(idx * 2)
    
    # Merge for card
    ws2.merge_cells(f"{col_letter}4:{next_col}4")
    ws2.merge_cells(f"{col_letter}5:{next_col}6")
    
    c_head = ws2[f"{col_letter}4"]
    c_head.value = kpi_title
    c_head.font = Font(name="Segoe UI", size=9, bold=True, color=color_hex)
    c_head.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
    c_head.alignment = Alignment(horizontal="center", vertical="center")
    
    c_val = ws2[f"{col_letter}5"]
    c_val.value = kpi_val
    c_val.font = Font(name="Segoe UI", size=18, bold=True, color=color_hex)
    c_val.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
    c_val.alignment = Alignment(horizontal="center", vertical="center")

# Detailed Phase Breakdown Table in Dashboard
ws2.cell(row=9, column=1, value="Thống Kê Chi Tiết Theo Giai Đoạn (Phase)").font = Font(name="Segoe UI", size=12, bold=True, color="0F172A")

phase_headers = ["STT", "Tên Giai Đoạn", "Số Lượng Task", "Tổng Số Ngày", "Số Task Hoàn Thành", "Tỷ Lệ Hoàn Thành"]
for col_i, ph in enumerate(phase_headers, start=1):
    cell = ws2.cell(row=10, column=col_i, value=ph)
    cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

phase_names = sorted(list(set([t[2] for t in tasks])))
for p_idx, p_name in enumerate(phase_names, start=11):
    p_tasks = [t for t in tasks if t[2] == p_name]
    p_count = len(p_tasks)
    p_days = sum([t[8] for t in p_tasks])
    p_done = len([t for t in p_tasks if t[9] == "Đã hoàn thành"])
    p_rate = p_done / p_count if p_count > 0 else 0
    
    r_data = [p_idx-10, p_name, p_count, p_days, p_done, p_rate]
    for c_i, val in enumerate(r_data, start=1):
        c = ws2.cell(row=p_idx, column=c_i, value=val)
        c.font = Font(name="Segoe UI", size=10)
        c.border = thin_border
        if c_i in [1, 3, 4, 5]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        elif c_i == 6:
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = '0.0%'
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")

# Auto-fit columns for both sheets
for ws in [ws1, ws2]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Avoid considering title merged cells for column width
            if cell.row in [1, 2, 3]:
                continue
            if cell.value:
                # Approximate string length for Vietnamese text
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Specific tweaks for sheet 1 column widths
ws1.column_dimensions['A'].width = 6
ws1.column_dimensions['B'].width = 15
ws1.column_dimensions['C'].width = 30
ws1.column_dimensions['D'].width = 42
ws1.column_dimensions['E'].width = 45
ws1.column_dimensions['F'].width = 16
ws1.column_dimensions['G'].width = 14
ws1.column_dimensions['H'].width = 14
ws1.column_dimensions['I'].width = 16
ws1.column_dimensions['J'].width = 18
ws1.column_dimensions['K'].width = 18
ws1.column_dimensions['L'].width = 12
ws1.column_dimensions['M'].width = 45

# Specific tweaks for sheet 2 column widths
ws2.column_dimensions['A'].width = 8
ws2.column_dimensions['B'].width = 35
ws2.column_dimensions['C'].width = 16
ws2.column_dimensions['D'].width = 16
ws2.column_dimensions['E'].width = 22
ws2.column_dimensions['F'].width = 20

output_path = r"d:\VTC_Academy\NihongoLife\NihongoLife\Bao_Cao\NihongoLife_Timeline_QuachThanhLong.xlsx"
wb.save(output_path)
print(f"Successfully generated Excel file at: {output_path}")
'''

with open(r"d:\VTC_Academy\NihongoLife\NihongoLife\Bao_Cao\generate_excel.py", "w", encoding="utf-8") as f:
    f.write(script_content)

print("Python generator written successfully.")
