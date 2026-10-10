from pathlib import Path
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Bao_Cao' / 'NihongoLife_Bao_Cao_Capstone_2026_Final.docx'
ASSET = ROOT / 'Bao_Cao'
READ = ROOT / 'docs' / 'readme'
TMP = ROOT / 'Bao_Cao' / '_report_assets'
TMP.mkdir(exist_ok=True)

NAVY = '17365D'; BLUE = '2F5597'; TEAL = '2E7D7B'; GOLD = 'C58B2A'; PALE = 'EAF1F8'; LIGHT = 'F4F6F9'; INK = '1F2937'; MUTED = '667085'; WHITE = 'FFFFFF'; RED = 'A61B1B'; GREEN = '287A4B'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths): row.cells[i].width = Cm(width)

def add_page_field(p):
    r = p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
    txt=OxmlElement('w:t'); txt.text='1'; end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for x in (begin,instr,sep,txt,end): r._r.append(x)

def add_toc(p):
    r=p.add_run(); b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin')
    i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text=' TOC \\o "1-3" \\h \\z \\u '
    s=OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'),'separate'); t=OxmlElement('w:t'); t.text='Mục lục sẽ tự cập nhật khi mở trong Microsoft Word.'
    e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end')
    for x in (b,i,s,t,e): r._r.append(x)

def keep_with_next(p): p.paragraph_format.keep_with_next = True

doc = Document()
sec = doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(3.0); sec.right_margin=Cm(2.0); sec.header_distance=Cm(1.1); sec.footer_distance=Cm(1.1)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(12); normal.font.color.rgb=RGBColor.from_string(INK)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; normal.paragraph_format.line_spacing=1.3; normal.paragraph_format.space_after=Pt(6)
for name,size,color,before,after in [('Title',28,NAVY,0,10),('Subtitle',14,MUTED,0,8),('Heading 1',17,NAVY,16,8),('Heading 2',14,BLUE,12,6),('Heading 3',12,TEAL,8,4)]:
    st=styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=(name!='Subtitle'); st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
styles['Caption'].font.name='Times New Roman'; styles['Caption'].font.size=Pt(10); styles['Caption'].font.italic=True; styles['Caption'].font.color.rgb=RGBColor.from_string(MUTED); styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; styles['Caption'].paragraph_format.space_after=Pt(8); styles['Caption']._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')

def header_footer(section):
    hp=section.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; hp.paragraph_format.space_after=Pt(2)
    rr=hp.add_run('NIHONGOLIFE  |  CAPSTONE REPORT'); rr.font.name='Arial'; rr.font.size=Pt(8); rr.font.bold=True; rr.font.color.rgb=RGBColor.from_string(MUTED)
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fp.paragraph_format.space_before=Pt(2)
    r=fp.add_run('VTC ACADEMY   •   '); r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED); add_page_field(fp)

def p(text='', style=None, bold_prefix=None, align=None):
    par=doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        par.add_run(bold_prefix).bold=True; par.add_run(text[len(bold_prefix):])
    else: par.add_run(text)
    if align is not None: par.alignment=align
    return par

def h(text, level=1): return p(text, f'Heading {level}')

def bullet(text):
    q=doc.add_paragraph(style='List Bullet'); q.paragraph_format.left_indent=Cm(0.75); q.paragraph_format.first_line_indent=Cm(-0.25); q.paragraph_format.space_after=Pt(3); q.add_run(text); return q

def table(headers, rows, widths=None, caption=None):
    if caption: cp=p(caption, 'Caption'); cp.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for j,x in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=x; set_cell_shading(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c)
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9.5); r.font.name='Arial'
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(t.rows[0])
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,x in enumerate(row):
            cells[j].text=str(x); cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[j]);
            if i%2: set_cell_shading(cells[j],LIGHT)
            for par in cells[j].paragraphs:
                par.paragraph_format.space_after=Pt(1); par.paragraph_format.line_spacing=1.05
                for r in par.runs: r.font.name='Times New Roman'; r.font.size=Pt(9.5)
    if widths: set_col_widths(t,widths)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

fig_no=0
def figure(path, caption, width=15.0):
    global fig_no
    path=Path(path)
    if not path.exists(): return
    fig_no+=1
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.keep_with_next=True; q.add_run().add_picture(str(path), width=Cm(width))
    cp=p(f'Hình {fig_no}. {caption}', 'Caption'); cp.paragraph_format.keep_with_next=False

def callout(title, text, color=PALE):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); set_cell_shading(c,color); set_cell_margins(c,140,180,140,180)
    c.text=''; pp=c.paragraphs[0]; pp.paragraph_format.space_after=Pt(0); a=pp.add_run(title+'  '); a.bold=True; a.font.color.rgb=RGBColor.from_string(NAVY); pp.add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Cover: editorial_cover
for _ in range(3): p('')
p('HỌC VIỆN CÔNG NGHỆ THÔNG TIN & THIẾT KẾ VTC ACADEMY', align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].bold=True
p('PROJECT 3 • CAPSTONE — LẬP TRÌNH GAME', align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.color.rgb=RGBColor.from_string(GOLD)
for _ in range(2): p('')
title=p('NIHONGOLIFE','Title',align=WD_ALIGN_PARAGRAPH.CENTER); title.runs[0].font.size=Pt(32)
p('ひばり町で、日本語と暮らそう',align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size=Pt(16)
p('GAME MÔ PHỎNG CUỘC SỐNG 3D KẾT HỢP HỌC TIẾNG NHẬT',align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].bold=True
figure(ASSET/'Logo'/'logo_dark.png','Nhận diện NihongoLife',7.0)
info=doc.add_table(rows=4,cols=2); info.alignment=WD_TABLE_ALIGNMENT.CENTER; info.style='Table Grid'; set_col_widths(info,[5.0,9.0])
for row,(a,b) in zip(info.rows,[('Học viên','Quách Thành Long'),('MSSV • Lớp','124010124034 • K24GD03'),('Giảng viên hướng dẫn','Nguyễn Ngọc Chấn'),('Chuyên ngành','Lập trình game')]):
    row.cells[0].text=a; row.cells[1].text=b
    for c in row.cells: set_cell_margins(c,120,160,120,160)
    row.cells[0].paragraphs[0].runs[0].bold=True; set_cell_shading(row.cells[0],PALE)
p('TP. Hồ Chí Minh, tháng 10 năm 2026',align=WD_ALIGN_PARAGRAPH.CENTER).paragraph_format.space_before=Pt(24)

# Main section
main=doc.add_section(WD_SECTION.NEW_PAGE); main.page_width=Cm(21); main.page_height=Cm(29.7); main.top_margin=Cm(2.2); main.bottom_margin=Cm(2.0); main.left_margin=Cm(3.0); main.right_margin=Cm(2.0); main.header_distance=Cm(1.1); main.footer_distance=Cm(1.1); main.header.is_linked_to_previous=False; main.footer.is_linked_to_previous=False; header_footer(main)

h('LỜI CẢM ƠN',1)
p('Em xin chân thành cảm ơn thầy Nguyễn Ngọc Chấn, giảng viên hướng dẫn, cùng các thầy cô VTC Academy đã định hướng phương pháp, góp ý về lập trình game, thiết kế trải nghiệm và quy trình triển khai đồ án. Em cảm ơn gia đình và bạn bè đã hỗ trợ trong suốt quá trình xây dựng NihongoLife. Báo cáo này được cập nhật theo trạng thái mã nguồn, tài liệu kỹ thuật và bằng chứng kiểm thử đến ngày 10/10/2026; những thành phần chưa được kiểm chứng thực tế được ghi rõ như giới hạn, không được trình bày như kết quả hoàn thành.')
p('Quách Thành Long',align=WD_ALIGN_PARAGRAPH.RIGHT).runs[0].bold=True

h('TÓM TẮT ĐỒ ÁN',1)
p('NihongoLife là game mô phỏng cuộc sống 3D hỗ trợ người Việt học tiếng Nhật ở mức nền tảng N5. Người chơi vào vai du học sinh tại Hibari-chō, học thông qua hội thoại và công việc đời sống: chào hỏi, mua sắm, đi tàu, gọi món, làm thêm, làm nông, chăm động vật và luyện thi. Thay vì tách bài học khỏi bối cảnh, hệ thống gắn mục tiêu ngôn ngữ với không gian, đối tượng, hành động và phản hồi ngay sau lựa chọn.')
p('Phiên bản hiện tại mở rộng qua bảy phase phát triển, từ nền tảng local và scenario engine đến UI/UX, tuyến tàu, Midori Island, hệ thống tiến trình - việc làm - kinh tế, IELTS/JLPT và lớp dịch vụ online có cơ chế offline fallback. Kiến trúc ưu tiên dữ liệu: scenario, progression, island catalog và exam package được tách khỏi mã điều khiển. Bằng chứng gần nhất gồm Stage B 19 Play Mode đạt, 1 test explicit bỏ qua, 16 EditMode đạt; IELTS có 5 EditMode và 1 PlayMode đạt; Midori Island có regression trọn vòng cùng đo hiệu năng trong Editor batchmode. Các thử nghiệm Supabase nhiều người, Agora hai máy và WebGL vẫn chưa hoàn tất.')
callout('Từ khóa', 'serious game; life simulation; contextual learning; Japanese N5; Unity; data-driven design; offline-first.')

h('MỤC LỤC',1); toc=doc.add_paragraph(); add_toc(toc); doc.add_page_break()

h('1. GIỚI THIỆU ĐỀ TÀI',1)
h('1.1. Bối cảnh và vấn đề',2)
p('Người mới học ngoại ngữ thường có thể nhận ra từ vựng trong bài tập nhưng phản ứng chậm khi gặp tình huống thực. Khoảng cách nằm ở việc kiến thức chưa gắn với người nói, mục đích, không gian và hệ quả. NihongoLife tiếp cận vấn đề bằng mô phỏng đời sống: mỗi bài học trở thành một việc cần hoàn thành trong thế giới 3D; lựa chọn sai tạo giải thích và đường thử lại thay vì chấm dứt trải nghiệm.')
h('1.2. Mục tiêu',2)
for x in ['Xây dựng vòng lặp khám phá - tương tác - phản hồi - tiến bộ có thể dùng lại cho nhiều bài học.','Tổ chức nội dung học theo dữ liệu để mở rộng mà không viết một gameplay riêng cho từng bài.','Kết nối tiến trình học với cấp độ, tiền Yen, việc làm, nhiệm vụ, nhu cầu sống và hoạt động trên đảo.','Hỗ trợ local/offline như đường chạy mặc định; các dịch vụ online là lớp tăng cường có trạng thái lỗi rõ ràng.','Đánh giá bằng kiểm thử thao tác thật trong Play Mode, kèm ảnh và log.']: bullet(x)
h('1.3. Phạm vi và giới hạn tuyên bố',2)
p('Báo cáo chỉ ghi nhận chức năng có mã nguồn, tài sản dữ liệu và/hoặc bằng chứng kiểm thử. Điểm JLPT/IELTS trong game là điểm luyện tập; estimated band không phải chứng chỉ. Nội dung Cambridge được giữ trong LocalContent đã gitignore và không phân phối trong build. Supabase có implementation và unit test merge nhưng chưa được xác nhận với tải nhiều người; Agora chưa có cuộc gọi hai máy; WebGL chưa có kết quả chạy build.')
figure(ASSET/'HinhAnh_Game'/'g01_gameplay.png','Gameplay tại Hibari-chō: HUD, người chơi và không gian tương tác.',15.0)

h('2. CƠ SỞ THIẾT KẾ VÀ GIÁ TRỊ GIÁO DỤC',1)
h('2.1. Học theo tình huống và phản hồi phục hồi',2)
p('Mỗi ScenarioDefinition kết hợp mục tiêu giao tiếp, learning target, node và lựa chọn. Node GoToArea, InspectItem và CollectItem nối câu chữ với hành động; DialogueChoice tác động điểm, mastery và story flags. Nhánh sai lý tưởng gồm ba bước: phản ứng phù hợp của NPC, giải thích ngắn và cơ hội lựa chọn lại. Cách này bảo toàn động lực và thể hiện chủ đề “nói sai cũng không sao”.')
h('2.2. Tham khảo thiết kế trải nghiệm',2)
p('The Sims 4 là tham chiếu cho cách tổ chức đời sống bằng nhu cầu, nhà ở, nghề nghiệp, địa điểm và các câu chuyện do người chơi tạo. NihongoLife không sao chép phạm vi sandbox, mà kế thừa nguyên tắc “việc thường ngày tạo thành mục tiêu chơi”: ăn, nghỉ, đi làm và di chuyển trở thành lý do tự nhiên để dùng ngôn ngữ. Sons of the Forest được dùng như tham chiếu tương phản cho tính trực tiếp của tương tác sinh tồn: trạng thái tài nguyên phải có phản hồi rõ, thao tác xây dựng/thu thập phải thấy kết quả trong thế giới, và vòng chuẩn bị - thực hiện - phục hồi cần dễ đọc. NihongoLife loại bỏ bạo lực, chỉ tiếp thu bài học về affordance, nhịp hành động và tính liên tục của trạng thái.')
h('2.3. Nguyên tắc đánh giá',2)
p('ScoringManager phân tách Vocabulary, Grammar, Listening, ResponseAccuracy và TaskCompletion. Mastery phản ánh lịch sử học, trong khi XP và level phản ánh tiến trình trò chơi; hai đại lượng không còn bị cộng chung. Sự tách biệt này tránh suy luận sai rằng chơi nhiều đồng nghĩa thành thạo ngôn ngữ. Muốn chứng minh hiệu quả giáo dục cần nghiên cứu người dùng trước/sau với công cụ đo độc lập.')

h('3. QUÁ TRÌNH PHÁT TRIỂN QUA 7 PHASE',1)
table(['Phase','Trọng tâm','Kết quả đã triển khai'],[
('1','Nền tảng local','Bootstrap, GameServices, điều khiển, save local, scenario engine, hội thoại và scoring.'),
('2','Vertical slice Hibari-chō','Thành phố, konbini, nhà hàng, phòng trọ, HUD, nhiệm vụ và 14 scenario N5.'),
('3','Mở rộng học tập và trình bày','JLPT/IELTS nền, trường học, Game Center/Kana Match, responsive UI, tài sản và pipeline.'),
('4','Di chuyển liên khu vực','Ga Hibari, vé, cổng soát, tàu tới Gakuen-mae/Minato và tuyến みどりじま; xử lý quay về.'),
('5','Midori Island','Scene 60_MidoriIsland, ruộng, cây theo thời gian UTC, vật nuôi, cửa hàng, sổ tay, vé về.'),
('6','Assessment và dịch vụ','IELTS Listening package local, grader/store/UI; Supabase/Gemini/Agora với mức kiểm chứng được công bố.'),
('7','Hoàn thiện life loop','HUD gọn, cursor/camera, Task Journal, level/XP, việc làm nhiều bước, kinh tế và merge progress.'),
],[2.0,4.2,10.2], 'Bảng 1. Tổng hợp bảy phase theo trạng thái repository ngày 10/10/2026.')
p('Các phase là lát cắt kỹ thuật của cùng một sản phẩm, không phải bảy game mode độc lập. Mỗi phase mở rộng vòng lặp cũ và phải dùng chung save, input, modal stack, HUD feed và service. Cách tổ chức này giảm nguy cơ một tính năng mới tạo ra hệ thống điểm hoặc giao diện không đồng bộ.')

h('4. GAMEPLAY, NHIỆM VỤ VÀ TIẾN TRÌNH',1)
h('4.1. Vòng lặp cốt lõi',2)
figure(READ/'d01_core_loop.png','Vòng lặp khám phá, mục tiêu, tương tác, phản hồi và tiến bộ.',14.5)
p('Người chơi nhận mục tiêu, tới đúng địa điểm, tương tác bằng F hoặc giao diện, thực hiện lựa chọn ngôn ngữ/hành động và nhận phản hồi. Kết quả cập nhật nhiệm vụ, mastery, XP, Yen hoặc vật phẩm tùy hoạt động. Khi UI dạng modal mở, gameplay input bị khóa thông qua UiModalStack; thao tác ESC đóng lớp trên cùng thay vì đồng thời kích hoạt nhiều cửa sổ.')
h('4.2. Sổ nhiệm vụ và catalog tiến trình',2)
p('Task Journal mở bằng N (J cũ vẫn tương thích), hợp nhất cốt truyện, việc làm, nông trại, học tập và nhiệm vụ hằng ngày. progression.json là nguồn cấu hình ngưỡng level, phần thưởng và mục tiêu. QuestService quản lý nhận, hủy, theo dõi và hoàn tất; HUD chỉ hiển thị một mục tiêu đang theo dõi để giảm nhiễu.')
figure(ASSET/'lifeloop-regression'/'04_journal_job.png','Sổ nhiệm vụ hiển thị nhóm việc làm và phần thưởng theo dữ liệu.',14.6)
h('4.3. Level, XP và tri thức',2)
p('progress.xp được chuẩn hóa thành tổng XP; level được suy ra từ bảng ngưỡng. Dữ liệu cũ được nâng cấp để không mất cấp. Điểm Knowledge được tách khỏi XP: Knowledge phục vụ điều kiện học và mastery, còn XP phản ánh độ tham gia/tiến trình trò chơi. Đây là ranh giới quan trọng giữa đánh giá giáo dục và phần thưởng gameplay.')

h('5. KINH TẾ VÀ VIỆC LÀM',1)
h('5.1. Mô hình dòng tiền',2)
p('Yen được dùng để mua thức ăn, vé tàu, hạt giống, dụng cụ và hàng hóa trên đảo. Nguồn thu hiện thực sự triển khai gồm bán nông sản và ba ca làm: konbini ¥600, sushi ¥700, Midori ¥500. Mỗi ca gồm nhiều bước và câu hỏi tiếng Nhật; trả lời sai không tính, hủy ca không trả lương, và một thời điểm chỉ có một ca hoạt động. Cơ chế này ngăn “nhấn nút nhận tiền” và gắn thu nhập với nhiệm vụ có thể quan sát.')
table(['Nguồn/chi','Điều kiện','Tác động'],[
('Konbini +¥600','Bê hàng, xếp 3 kệ, chỉ đường 2 khách, tính tiền, báo Ito','Thu nhập + thực hành phục vụ'),
('Sushi +¥700','Nhận 2 order, lấy đúng món, giao đúng bàn, báo Aoki','Thu nhập + nghe/đọc order'),
('Midori +¥500','Nhận hạt, xới, gieo, tưới, cho thú ăn, báo Hana','Thu nhập + từ vựng nông trại'),
('Bán nông sản','Có vật phẩm thu hoạch trong balo','Chuyển sản phẩm thành Yen'),
('Chi tiêu','Mua đồ, hạt/dụng cụ, vé tàu','Tạo quyết định và vòng tái đầu tư'),
],[4.0,7.2,5.2], 'Bảng 2. Các dòng tiền đã triển khai.')
figure(ASSET/'lifeloop-regression'/'06_carrying_box.png','Ca làm tại konbini yêu cầu thao tác vật lý và chuỗi mục tiêu.',14.6)
figure(ASSET/'lifeloop-regression'/'20_sushi_order.png','Ca làm sushi: đọc order tiếng Nhật và giao đúng món.',14.6)

h('6. MIDORI ISLAND: NÔNG TRẠI VÀ VẬT NUÔI',1)
h('6.1. Cấu trúc khu vực',2)
p('Scene 60_MidoriIsland đã được duyệt và dựng bằng IslandBuilder.Build, không dùng MenuItem. Khu vực gồm ga, quảng trường, sáu ô ruộng, chuồng với bò, alpaca, lừa, ngựa, chó Shiba, cửa hàng, điểm ngắm cảnh và biển học từ. Bờ biển có giới hạn va chạm và cơ chế đưa người chơi về ga khi rơi khỏi đảo. Khi đảo nạp, mặt trời của thành phố bị tắt để tránh hai nguồn sáng định hướng.')
figure(ASSET/'island-regression'/'18_island_overview.png','Toàn cảnh Midori Island với nông trại, chuồng nuôi và ga.',15.0)
h('6.2. Quy trình farming',2)
p('Dữ liệu cây trồng nằm trong island_catalog.json: giá, thời gian lớn, sản lượng và từ vựng JA/EN/VI. Trạng thái ruộng lưu trong PlayerProgressDto.island. Cây lớn theo dấu thời gian UTC nên tiếp tục phát triển khi người chơi rời đảo. Thao tác xới có thời lượng theo công cụ; không có bình tưới thì không thể tưới. Chu trình kiểm thử đã đi qua xới đất, chọn cà rốt, tưới ba giai đoạn, thu hoạch hai củ và bán.')
table(['Bước','Điều kiện','Trạng thái/đầu ra'],[
('1. Chuẩn bị','Ô trống; có tay/xẻng/cuốc','Xới 7,0 s / 2,6 s / 1,6 s'),('2. Gieo','Có hạt giống','Ô ghi loại cây và thời điểm'),('3. Tưới','Có bình tưới','Bắt đầu/tiếp tục tăng trưởng'),('4. Chờ','Đủ mốc UTC','Chuyển model qua các stage'),('5. Thu hoạch','Phase Ready','Nông sản vào balo, mở thành tích'),('6. Bán/tái đầu tư','Có nông sản','Nhận Yen, mua hạt/dụng cụ mới')
],[3.0,6.2,7.2], 'Bảng 3. Quy trình farming đã triển khai.')
figure(ASSET/'island-regression'/'08_farm_seed_picker.png','Chọn hạt giống tại ô ruộng.',14.6)
figure(ASSET/'island-regression'/'11_farm_ready.png','Cây đạt trạng thái sẵn sàng thu hoạch.',14.6)
figure(ASSET/'island-regression'/'12_harvested.png','Thu hoạch đưa nông sản vào balo.',14.6)
h('6.3. Chăn nuôi và học từ vựng',2)
p('IslandAnimal hiển thị thẻ tương tác, cho ăn, vuốt ve và phản hồi quan hệ. Catalog quy định thức ăn và tên đa ngôn ngữ. Regression ghi nhận tương tác với bò, cho ăn hai lần, vuốt ve, động vật di chuyển tối đa 3,49 m trong 4 giây; sổ tay ghi 13 từ và thành tích first_harvest, first_sale. Hệ thống hiện là chăm sóc và làm quen, chưa phải breeding hoặc sản xuất sữa/lông.')
figure(ASSET/'island-regression'/'16_animal_fed.png','Cho vật nuôi ăn và cập nhật trạng thái quan hệ.',14.6)
h('6.4. Tuyến tàu và chống soft-lock',2)
p('Ga みどりじま có giá vé ¥450 từ catalog. Chiều về đi qua IslandTicketKiosk, IslandTrainDoor, màn hình di chuyển và trở lại 20_StationDistrict. Hệ thống tránh thu tiền hai lần. Nếu người chơi hết tiền và không có vật phẩm để bán, game cấp vé hỗ trợ để không bị mắc kẹt trên đảo.')

h('7. HỆ THỐNG HỌC TẬP VÀ KHẢO THÍ',1)
h('7.1. Scenario N5 và hội thoại',2)
p('Repository có 14 ScenarioDefinition phục vụ tuyến truyện từ ngày đầu tới lễ hội. Engine hỗ trợ node hội thoại, đi tới khu vực, xem/nhặt vật phẩm, nhánh, hoàn thành và thất bại. setFlags và điều kiện done:scenarioId tạo bộ nhớ truyện. DialogueView trình bày tiếng Nhật, cách đọc, romaji và nghĩa theo chế độ Guided/Practice/Assessment.')
figure(ASSET/'HinhAnh_Game'/'g09_dialogue_choice.png','Hội thoại có lựa chọn và phản hồi trong ngữ cảnh.',14.5)
h('7.2. JLPT',2)
p('ExamRepository nạp đề JLPT N5 rút gọn tự soạn. ExamManager quản lý attempt; ExamPlayUI hiển thị câu hỏi, thời gian, điều hướng và nộp phần; ExamAttemptRecord lưu lịch sử. Nhãn UI ghi rõ đây là đề luyện tập, không phải đề chính thức. Kết quả trong game dùng cho tự đánh giá và phản hồi của cô Morita.')
h('7.3. IELTS Listening',2)
p('Phase 6 bổ sung schema nihongolife.ielts.v1, IeltsLibrary, IeltsGrader, IeltsAttemptStore và IeltsTestUI. Chế độ luyện tập cho phép phát/dừng/tua; chế độ thi phát mỗi phần một lần, không dừng, tự chuyển phần và tự nộp sau thời gian kiểm tra. UI có form điền, điều hướng 1-40, tự lưu, tiếp tục và review. Gói Cambridge 14 Test 1 chỉ nằm trong LocalContent gitignore để tôn trọng bản quyền; test E2E tự bỏ qua nếu máy không có gói local.')
table(['Bằng chứng','Kết quả','Phạm vi'],[('IeltsGraderTests','5/5 đạt','Dữ liệu giả lập: optional text, word limit, either order, scoring'),('IeltsListeningPlayModeTests','1/1 đạt','Luồng UI E2E; phụ thuộc local package'),('Stage B EditMode','16/16 đạt','Bao gồm test dữ liệu/engine mới'),('Stage B PlayMode','19 đạt, 1 explicit bỏ qua','Snapshot tích hợp ngày 10/10/2026')],[4.8,4.0,7.6], 'Bảng 4. Bằng chứng kiểm thử mới nhất.')

h('8. KIẾN TRÚC HỆ THỐNG VÀ DỮ LIỆU',1)
figure(READ/'d08_architecture.png','Kiến trúc phân lớp UI - Gameplay - Service - Data - hạ tầng tùy chọn.',15.0)
h('8.1. Bootstrap và service',2)
p('00_Bootstrap tạo AppRoot và đăng ký các service dùng lâu dài qua GameServices. Các scene zone chỉ chứa đối tượng theo khu vực; UI và runtime component đăng ký/hủy event theo vòng đời. SceneFlowController điều phối nạp additive, spawn và hội thoại khi di chuyển. Cấu trúc này tránh nhân bản service và giảm phụ thuộc trực tiếp giữa gameplay với backend.')
h('8.2. Mô hình dữ liệu',2)
figure(READ/'d18_erd.png','ERD Supabase và quan hệ dữ liệu online.',14.8)
p('PlayerProgressDto là snapshot local gồm campaign, inventory, exam attempts, XP, island và quests. SupabaseProgressRepository lưu progress_json và ghép dữ liệu khi đăng nhập. Lỗi từng làm rơi inventory, island và quests trong merge đã được sửa và có ProgressMergeTests; tuy nhiên kết nối Supabase thật chưa được chạy lại cho phase mới. island_catalog.json và progression.json là catalog cấu hình, giúp giá và ngưỡng không bị viết cứng trong script.')
h('8.3. Online/offline',2)
table(['Thành phần','Đã triển khai','Chưa được khẳng định'],[
('Local/offline','Save local, scenario, farming, việc làm, thi rút gọn, fallback UI','Không cần tài khoản để chơi vòng cốt lõi'),
('Supabase','Auth/repository/presence/chat/friends/co-op code; merge unit test','Tải nhiều người và merge phase mới trên backend thật'),
('Gemini','Pronunciation/NPC/Writing integration; model gemini-flash-latest từng trả 200','Độ ổn định, chi phí và chất lượng chấm quy mô lớn'),
('Agora','SDK và luồng lớp học F8','Cuộc gọi thật giữa hai máy'),
('WebGL','Có lưu ý pointer lock sau click canvas','Chưa chạy build WebGL hiện tại')
],[3.2,7.1,6.1], 'Bảng 5. Ranh giới xác minh online/offline.')

h('9. PHÂN TÍCH UI/UX',1)
h('9.1. HUD và phân cấp thông tin',2)
p('HUD mới bỏ dải phím tắt cố định trên desktop, giữ portrait, level, Yen và năm vòng nhu cầu ở góc dưới trái. Cảnh báo chỉ xuất hiện dưới 20%; góc trên trái giữ một mục tiêu; thông báo được hợp nhất qua HudFeed; prompt F ưu tiên hơn lời nhắc R. HUD ẩn khi thi hoặc mini-game. Bộ ảnh regression kiểm tra 1920×1080, 1600×900, 1366×768 và 1280×720.')
figure(ASSET/'lifeloop-regression'/'01_hud_1920x1080.png','HUD gọn ở 1920×1080: trạng thái, mục tiêu và prompt không chồng lấn.',15.0)
h('9.2. Input, camera và modal',2)
p('CursorDirector là nguồn quyết định con trỏ. Mặc định game khóa chuột để xoay camera; giữ Ctrl giải phóng con trỏ; modal tự mở con trỏ; mất focus sẽ nhả. Người chơi vẫn có thể chọn click-to-move kiểu life-sim trong Settings. UiModalStack bảo đảm ESC đóng cửa sổ trên cùng; bài thi hỏi xác nhận trước khi rời. Thiết kế này giảm xung đột từng xảy ra khi 13 script cùng đọc ESC.')
h('9.3. Khả năng tiếp cận và bản địa hóa',2)
p('UI hỗ trợ Việt/Anh/Nhật; Noto Sans JP dùng cho kana/kanji và dấu Việt. Giao diện cần tiếp tục được kiểm tra với chuỗi dài, font fallback, tương phản và người dùng thật. Midori cho đổi ngôn ngữ học JA/EN, nhưng ngôn ngữ giao diện và ngôn ngữ mục tiêu được coi là hai thiết lập khác nhau.')

h('10. KIỂM THỬ VÀ ĐÁNH GIÁ',1)
h('10.1. Chiến lược',2)
p('EditMode dùng để kiểm tra catalog, graph, grader và merge; PlayMode điều khiển nhân vật/UI theo đường input thật, chụp ảnh từng bước và kiểm tra collider/spawn. Không coi file script tồn tại là bằng chứng tính năng. Ảnh regression được xem cùng log và XML test-run; test explicit được ghi riêng, không tính như test đã chạy.')
table(['Snapshot','Total','Passed','Failed','Skipped/Explicit'],[
('Final 08/10 PlayMode','17','16','0','1'),('Final 08/10 EditMode','11','11','0','0'),('Stage A 09/10 PlayMode','19','18','0','1'),('Stage A 09/10 EditMode','11','11','0','0'),('Stage B 10/10 PlayMode','20','19','0','1'),('Stage B 10/10 EditMode','16','16','0','0'),('IELTS riêng','6','6','0','0')
],[5.0,2.6,2.6,2.6,3.6], 'Bảng 6. Kết quả test-run có XML/log trong Bao_Cao.')
h('10.2. Regression Midori Island',2)
p('Hai test IslandPlayModeTests đi qua luồng chức năng và kiểm tra không gian. Log chức năng ghi tổng 78,2 giây cho vòng đi đảo, mua, farming, bán, chăm bò và quay về. Log hiệu năng Editor batchmode ghi 2.686 frame, trung bình 1,5 ms, p95 3,1 ms, cực đại 7,2 ms; 248 renderer, 742 batch, 67 SetPass, khoảng 600.878 tam giác. Đây là số đo Editor trong bài test, không phải benchmark build/player hay thiết bị mục tiêu.')
figure(ASSET/'island-regression'/'30_perf_view.png','Khung hình dùng trong kiểm tra hiệu năng Midori Island.',14.6)
h('10.3. Rủi ro còn lại',2)
for x in ['Supabase nhiều người, xung đột sync thực tế và co-op chưa có test tải.','Agora chưa có cuộc gọi hai thiết bị; microphone thật chưa được xác nhận toàn bộ.','WebGL chưa chạy; pointer lock phụ thuộc click vào canvas.','Nội dung tiếng Nhật cần giáo viên/người bản ngữ duyệt; điểm AI chỉ mang tính tham khảo.','Một số tài sản cửa hàng/công nghệ mới là đồ sưu tầm, chưa đặt được vào phòng; tab thời trang Midori cố ý để trống và giải thích trung thực.']: bullet(x)

h('11. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN',1)
p('NihongoLife đã hình thành một vertical slice có chiều rộng đáng kể: thế giới nhiều khu vực, 14 scenario N5, hoạt động đời sống, luyện thi, mini-game, tuyến tàu, nông trại, vật nuôi, việc làm, kinh tế và progression. Điểm mạnh kỹ thuật không nằm ở số lượng màn hình riêng lẻ mà ở khả năng dùng chung service, save, input, modal và catalog dữ liệu. Điểm mạnh giáo dục là đặt ngôn ngữ vào hành động có mục đích và cho phép phục hồi sau lỗi.')
p('Ưu tiên tiếp theo là hoàn tất kiểm thử dịch vụ thật và build WebGL; review ngôn ngữ; quan sát người học để đánh giá tải nhận thức, khả năng tìm mục tiêu và chất lượng feedback; cân bằng vòng Yen - việc làm - farming dựa trên telemetry thay vì ước lượng. Các mở rộng như breeding, đặt đồ công nghệ, thời trang hoặc ngân hàng đề lớn chỉ nên được đưa vào báo cáo khi có implementation và bằng chứng tương ứng.')

h('12. TÀI LIỆU THAM KHẢO',1)
refs=[
('[1]','Electronic Arts. The Sims 4 - Official Game Overview. https://www.ea.com/games/the-sims/the-sims-4/'),
('[2]','Endnight Games / Newnight. Sons Of The Forest - Steam Store. https://store.steampowered.com/app/1326470/Sons_Of_The_Forest/'),
('[3]','Unity Technologies. Unity 6 Manual and Scripting API. https://docs.unity3d.com/'),
('[4]','The Japan Foundation. JLPT N5: summary of linguistic competence and test sections. https://www.jlpt.jp/e/about/levelsummary.html'),
('[5]','The Japan Foundation. Irodori: Japanese for Life in Japan. https://www.irodori.jpf.go.jp/'),
('[6]','NihongoLife repository. README.md; Docs/ARCHITECTURE.md; STORY_BIBLE.md; SCENARIO_SYSTEM.md; EXAM_SYSTEM.md; PROGRESSION_AND_EXAM_INTEGRITY.md; TEAM_TASKS.md.'),
('[7]','NihongoLife test evidence. Bao_Cao/stageA_*.xml, stageB_*.xml, ielts_*.xml, island-regression/*.txt và ảnh regression, snapshot 10/10/2026.')]
for a,b in refs:
    q=doc.add_paragraph(); q.paragraph_format.first_line_indent=Cm(-0.8); q.paragraph_format.left_indent=Cm(0.8); q.add_run(a+' ').bold=True; q.add_run(b)

h('PHỤ LỤC A. BẰNG CHỨNG HÌNH ẢNH',1)
figure(ASSET/'island-regression'/'01_onboard_to_midori.png','Chọn tuyến tàu tới みどりじま.',14.6)
figure(ASSET/'island-regression'/'03_shop_1920x1080.png','Cửa hàng đảo ở độ phân giải 1920×1080.',14.6)
figure(ASSET/'lifeloop-regression'/'31_farm_paid.png','Hoàn tất ca làm nông và nhận lương.',14.6)
figure(ASSET/'island-regression'/'21_back_at_hibari.png','Trở lại ga Hibari sau chuyến đi đảo.',14.6)

# update fields on open
settings=doc.settings._element; upd=OxmlElement('w:updateFields'); upd.set(qn('w:val'),'true'); settings.append(upd)

# Core metadata
doc.core_properties.title='NihongoLife - Báo cáo đồ án Project 3 Capstone'
doc.core_properties.subject='Game mô phỏng cuộc sống 3D kết hợp học tiếng Nhật'
doc.core_properties.author='Quách Thành Long'
doc.core_properties.keywords='NihongoLife, Unity, Japanese learning, Capstone'
doc.save(OUT)
print(OUT)
