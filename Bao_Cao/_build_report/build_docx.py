import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

import extra_content as X
import extra_r7 as R7

P = Path(__file__).parent
OUT = Path(r"D:/VTC_Academy/NihongoLife/NihongoLife/Bao_Cao/NihongoLife_Bao_Cao_Capstone.docx")
data = json.loads((P / "report-content.json").read_text(encoding="utf-8"))

INK, GOLD, RED, MUTED, TEXT, SOFT, LINE = "14213D", "C9901F", "C8463D", "5B6B82", "22304A", "F2F5FA", "D5DCE8"
HEAD_FONT, BODY_FONT, JP_FONT = "Segoe UI Semibold", "Calibri", "Yu Gothic"
CONTENT_W = 9020  # twips: A4 21cm - 2.5 - 2.0 = 16.5cm ≈ 9354; leave a little air

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0)
sec.header_distance = sec.footer_distance = Cm(1.0)


def set_font(style_or_run, name, size=None, color=None, bold=None, east=JP_FONT):
    f = style_or_run.font
    f.name = name
    if size: f.size = Pt(size)
    if color: f.color.rgb = RGBColor.from_string(color)
    if bold is not None: f.bold = bold
    el = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    rpr = el.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(k), name)
    fonts.set(qn("w:eastAsia"), east)


def style(name, font, size, color, before=0, after=6, line=1.15, bold=None, italic=None, align=None, base=None):
    st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
    if base: st.base_style = doc.styles[base]
    set_font(st, font, size, color, bold)
    if italic is not None: st.font.italic = italic
    pf = st.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if align is not None: pf.alignment = align
    return st


style("Normal", BODY_FONT, 11.5, TEXT, 0, 7, 1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
h1 = style("Heading 1", HEAD_FONT, 24, INK, 0, 4, 1.05, bold=False)
h2 = style("Heading 2", HEAD_FONT, 14, RED, 14, 5, 1.1, bold=False)
h3 = style("Heading 3", HEAD_FONT, 12, INK, 10, 4, 1.1, bold=False)
for h in (h1, h2, h3):
    h.paragraph_format.keep_with_next = True
    h.font.italic = False
style("Kicker", HEAD_FONT, 10, GOLD, 0, 2, 1.0, base="Normal", align=WD_ALIGN_PARAGRAPH.LEFT)
style("Lead", BODY_FONT, 13, MUTED, 0, 16, 1.2, italic=True, base="Normal", align=WD_ALIGN_PARAGRAPH.LEFT)
style("Caption", BODY_FONT, 9.5, MUTED, 4, 14, 1.1, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
style("Cell", BODY_FONT, 10, TEXT, 2, 2, 1.1, base="Normal", align=WD_ALIGN_PARAGRAPH.LEFT)
style("CellHead", HEAD_FONT, 10, "FFFFFF", 2, 2, 1.1, base="Normal", align=WD_ALIGN_PARAGRAPH.LEFT)
style("Note", BODY_FONT, 10.5, INK, 0, 0, 1.25, base="Normal", align=WD_ALIGN_PARAGRAPH.LEFT)
for n in ("TOC 1", "TOC 2"):
    st = style(n, BODY_FONT, 11 if n == "TOC 1" else 10.5, INK if n == "TOC 1" else TEXT, 4 if n == "TOC 1" else 0, 2, 1.1, base="Normal",
               align=WD_ALIGN_PARAGRAPH.LEFT)
    if n == "TOC 1": st.font.bold = True
    else: st.paragraph_format.left_indent = Cm(0.6)
style("Table of Figures", BODY_FONT, 10.5, TEXT, 0, 3, 1.1, base="Normal", align=WD_ALIGN_PARAGRAPH.LEFT)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    e = OxmlElement("w:shd"); e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), fill); tcpr.append(e)


def cell_margins(tbl, top=90, bottom=90, left=140, right=140):
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tbl._tbl.tblPr.append(m)


def borders(tbl, color=LINE, inside=True, size=4):
    b = OxmlElement("w:tblBorders")
    keys = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for k in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement(f"w:{k}")
        if k in keys:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tbl._tbl.tblPr.append(b)


def fixed_table(rows, cols, widths):
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    w = pr.find(qn("w:tblW")); w.set(qn("w:type"), "dxa"); w.set(qn("w:w"), str(sum(widths)))
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); pr.append(lay)
    grid = t._tbl.tblGrid
    for g in list(grid): grid.remove(g)
    for wd in widths:
        g = OxmlElement("w:gridCol"); g.set(qn("w:w"), str(wd)); grid.append(g)
    for r in t.rows:
        for j, c in enumerate(r.cells):
            c.width = Inches(widths[j] / 1440)
    return t


def scale_widths(widths, total=CONTENT_W):
    s = sum(widths)
    out = [round(w * total / s) for w in widths]
    out[-1] = total - sum(out[:-1])
    return out


def table(headers, rows, widths=None, first_bold=True):
    widths = scale_widths(widths or [1] * len(headers))
    t = fixed_table(len(rows) + 1, len(headers), widths)
    borders(t, LINE, True)
    cell_margins(t)
    for i, vals in enumerate([headers] + rows):
        tr = t.rows[i]._tr.get_or_add_trPr()
        tr.append(OxmlElement("w:cantSplit"))
        if i == 0: tr.append(OxmlElement("w:tblHeader"))
        for j, v in enumerate(vals):
            c = t.rows[i].cells[j]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.style = doc.styles["CellHead" if i == 0 else "Cell"]
            r = p.add_run(v)
            if i == 0:
                shade(c, INK)
            else:
                if i % 2 == 0: shade(c, "F7F9FC")
                if j == 0 and first_bold: r.bold = True; r.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(title, text):
    t = fixed_table(1, 1, [CONTENT_W])
    borders(t, "F2D493", False, 6)
    cell_margins(t, 160, 160, 220, 220)
    c = t.rows[0].cells[0]; shade(c, "FFF8E6")
    p = c.paragraphs[0]; p.style = doc.styles["Note"]
    r = p.add_run(title + "  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(GOLD)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def field(par, instr, placeholder=""):
    r = par.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
    r = par.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate"); r._r.append(f2)
    par.add_run(placeholder)
    r = par.add_run(); f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end"); r._r.append(f3)


chapter_no = [0]


def _jpg(path):
    q = Path(path)
    if q.suffix == ".png" and q.parts[0] in ("shots", "shots7") and (P / q.with_suffix(".jpg")).exists() and "vtc-logo" not in q.name:
        return str(q.with_suffix(".jpg"))
    return path


def figure(path, caption, width):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(P / _jpg(path)), width=Inches(width))
    cp = doc.add_paragraph(style="Caption")
    r = cp.add_run("Hình "); r.bold = True; r.italic = False; r.font.color.rgb = RGBColor.from_string(INK)
    pre = cp.add_run(f"{chapter_no[0]}." if chapter_no[0] else "")
    pre.bold = True; pre.italic = False; pre.font.color.rgb = RGBColor.from_string(INK)
    field(cp, "SEQ Hình \\* ARABIC" + (" \\s 1" if chapter_no[0] else ""), "1")
    for run in cp.runs[2:]:
        run.bold = True; run.italic = False; run.font.color.rgb = RGBColor.from_string(INK)
    cp.add_run(f" — {caption}")


def para(text, st=None):
    p = doc.add_paragraph(style=st)
    if _pb[0]:
        p.paragraph_format.page_break_before = True; _pb[0] = False
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        if i: p.add_run().add_break()
        p.add_run(ln)
    if len(lines) > 1: p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


_pb = [False]


def fig_cell(cell, path, caption, width):
    pp = cell.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(str(P / _jpg(path)), width=Inches(width))
    cp = cell.add_paragraph(style="Caption"); cp.paragraph_format.space_after = Pt(4)
    r = cp.add_run("Hình "); r.bold = True; r.italic = False; r.font.color.rgb = RGBColor.from_string(INK)
    pre = cp.add_run(f"{chapter_no[0]}." if chapter_no[0] else ""); pre.bold = True; pre.italic = False; pre.font.color.rgb = RGBColor.from_string(INK)
    n0 = len(cp.runs)
    field(cp, "SEQ Hình \\* ARABIC" + (" \\s 1" if chapter_no[0] else ""), "1")
    for run in cp.runs[n0:]:
        run.bold = True; run.italic = False; run.font.color.rgb = RGBColor.from_string(INK)
    cp.add_run(f" — {caption}")


def fig_pair(items):
    half = (CONTENT_W - 120) // 2
    t = fixed_table(1, 2, [half + 60, half + 60]); borders(t, "FFFFFF", False); cell_margins(t, 40, 40, 40, 40)
    for k, (path, cap) in enumerate(items):
        fig_cell(t.rows[0].cells[k], path, cap, half / 1440)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def renumber(text):
    m = re.match(r"^(\d+)\.(\d+)\.(.*)$", text)
    if m and chapter_no[0]:
        return f"{chapter_no[0]}.{m.group(2)}.{m.group(3)}"
    return text


def render_part(part):
    kind = part[0]
    if kind == "h":
        doc.add_paragraph(renumber(part[1]), style="Heading 2")
    elif kind == "h3":
        doc.add_paragraph(part[1], style="Heading 3")
    elif kind == "p":
        para(part[1])
    elif kind == "fig":
        figure(part[1], part[2], part[3])
    elif kind == "pair":
        fig_pair(part[1])
    elif kind == "table":
        table(part[1], part[2], part[3] if len(part) > 3 else None)
    elif kind == "callout":
        callout(part[1], part[2])


def page_break():
    _pb[0] = True


# ─────────── Cover ───────────
sec.different_first_page_header_footer = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
p.add_run().add_picture(str(P / "shots/vtc-logo.png"), width=Cm(7.5))
p = para("HỌC VIỆN CÔNG NGHỆ THÔNG TIN & THIẾT KẾ VTC ACADEMY", "Kicker"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.color.rgb = RGBColor.from_string(MUTED); p.paragraph_format.space_after = Pt(18)

band = fixed_table(1, 1, [CONTENT_W]); borders(band, INK, False); cell_margins(band, 380, 380, 400, 400)
c = band.rows[0].cells[0]; shade(c, INK)
cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cp.add_run("BÁO CÁO ĐỒ ÁN  ·  PROJECT 3 CAPSTONE"); set_font(r, HEAD_FONT, 11, "F2B233")
r.font.spacing = Pt(1.5) if hasattr(r.font, "spacing") else None
for text, fnt, size, col, after in [("NIHONGOLIFE", HEAD_FONT, 46, "FFFFFF", 0), ("ひばり町で、日本語と暮らそう", JP_FONT, 16, "F2B233", 10),
                                     ("Game mô phỏng cuộc sống 3D kết hợp học tiếng Nhật N5", BODY_FONT, 14, "DCE3EE", 0)]:
    q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after = Pt(after); q.paragraph_format.space_before = Pt(4)
    rr = q.add_run(text); set_font(rr, fnt, size, col, east=JP_FONT if fnt != JP_FONT else JP_FONT)
    if fnt == JP_FONT: set_font(rr, JP_FONT, size, col)
doc.add_paragraph().paragraph_format.space_after = Pt(4)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(14)
p.add_run().add_picture(str(P / _jpg("shots/01_menu.png")), width=Cm(16.5))

info = [("Ngành", "Công nghệ thông tin"), ("Chuyên ngành", "Lập trình game"), ("Giảng viên hướng dẫn", "Nguyễn Ngọc Chấn"),
        ("Học viên", "Quách Thành Long"), ("MSSV", "124010124034"), ("Lớp", "K24GD03")]
it = fixed_table(3, 4, scale_widths([1900, 2700, 1900, 2520])); borders(it, LINE, True); cell_margins(it, 70, 70, 140, 140)
for k, (lab, val) in enumerate(info):
    rr_, cc_ = k // 2, (k % 2) * 2
    a, b = it.rows[rr_].cells[cc_], it.rows[rr_].cells[cc_ + 1]
    shade(a, SOFT)
    pa = a.paragraphs[0]; pa.style = doc.styles["Cell"]; x = pa.add_run(lab); x.font.color.rgb = RGBColor.from_string(MUTED)
    pb = b.paragraphs[0]; pb.style = doc.styles["Cell"]; y = pb.add_run(val); y.bold = True; y.font.color.rgb = RGBColor.from_string(INK)
p = para("TP. Hồ Chí Minh, tháng 10 năm 2026", "Caption"); p.paragraph_format.space_before = Pt(14)

# ─────────── Header / footer ───────────
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run("NihongoLife  ·  Báo cáo Project 3 Capstone"); set_font(r, BODY_FONT, 8.5, MUTED)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Quách Thành Long  ·  K24GD03  ·  Trang "); set_font(r, BODY_FONT, 8.5, MUTED)
field(fp, "PAGE", "1")
for run in fp.runs[1:]: set_font(run, BODY_FONT, 8.5, INK)

# ─────────── TOC + list of figures ───────────
page_break()
para("MỤC LỤC", "Kicker")
doc.add_paragraph("Nội dung báo cáo", style="Heading 1").style = doc.styles["Heading 1"]
toc_title = doc.paragraphs[-1]
# keep the TOC page title out of the TOC itself
toc_title.style = doc.styles["Normal"]; set_font(toc_title.runs[0], HEAD_FONT, 24, INK); toc_title.paragraph_format.space_after = Pt(14)
p = doc.add_paragraph(); field(p, 'TOC \\o "1-1" \\h \\z \\u', "Mở bằng Word và cập nhật mục lục (F9).")
page_break()
para("DANH MỤC", "Kicker")
q = para("Hình ảnh và sơ đồ"); set_font(q.runs[0], HEAD_FONT, 24, INK); q.paragraph_format.space_after = Pt(14)
p = doc.add_paragraph(); field(p, 'TOC \\h \\z \\c "Hình"', "Cập nhật trường để hiện danh mục hình.")

# ─────────── Chapters ───────────
X.FIGS.update(R7.FIGS_OVERRIDE)
for k, v in R7.PATCH_EXTRA.items():
    X.PATCH.setdefault(k, {}).update(v)
R7.APPEND.setdefault(23, [])[:0] = [("h", "22.4. Danh sách bộ kiểm thử Play Mode"), ("table",) + R7.TEST_TABLE]
X.TABLE_REPLACE[23][0] = R7.RESULTS_TABLE
X.BUGS[2].extend(R7.BUGS_EXTRA)
items = []
for n0, it in enumerate(data):
    it["_orig"] = n0
    items.append(it)
    for after, new in R7.NEW_CHAPTERS:
        if after == n0:
            new = dict(new); new["_orig"] = None; items.append(new)
running = [0]
for idx, item in enumerate(items):
    n = item["_orig"] if item["_orig"] is not None else -1
    page_break()
    title = item["title"].replace("Tài liệu đối chiếu", "Tài liệu tham khảo")
    item["lead"] = {23: "Kết quả kiểm thử, các lỗi đã sửa và hạng mục còn theo dõi.",
                    26: "Tài liệu thiết kế của dự án và nguồn tham khảo bên ngoài."}.get(n, item["lead"])
    m = re.match(r"^(\d+)\.\s*(.*)$", title)
    if m:
        running[0] += 1
        chapter_no[0] = running[0]
        para(f"CHƯƠNG {chapter_no[0]:02d}", "Kicker")
        doc.add_paragraph(f"{chapter_no[0]}. {m.group(2)}", style="Heading 1")
    else:
        chapter_no[0] = 0
        para("NIHONGOLIFE", "Kicker")
        doc.add_paragraph(title, style="Heading 1")
    lead = para(item["lead"], "Lead")
    pb = lead._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "6"); bt.set(qn("w:space"), "8"); bt.set(qn("w:color"), LINE); bdr.append(bt); pb.append(bdr)

    patches = X.PATCH.get(n, {})
    tables_seen = 0
    skip = False
    drop = X.DROP_SUBSECTIONS.get(n, [])
    figs = list(X.FIGS.get(n, []))
    fig_placed = False
    for part in item["parts"]:
        kind = part[0]
        if kind == "h":
            skip = any(part[1].startswith(d) for d in drop)
            if n == 23 and part[1].startswith("22.1"):
                head, cols, rows = X.BUGS
                doc.add_paragraph(renumber(head), style="Heading 2")
                table(cols, rows, [2600, 3400, 3400])
                skip = True
                continue
            if skip: continue
            # place the first screenshot right before the first sub-heading of UI chapters
            if figs and not fig_placed and n in X.FIGS and figs[0][0].startswith("shots"):
                figure(*figs.pop(0)); fig_placed = True
            doc.add_paragraph(renumber(part[1]), style="Heading 2")
        elif skip:
            continue
        elif kind == "p":
            txt = part[1]
            for old, new in patches.items():
                if txt.startswith(old):
                    txt = new
                    break
            if txt is None: continue
            if n == 0 and txt.startswith("TP. Hồ Chí Minh"):
                p = para("TP. Hồ Chí Minh, ngày 08 tháng 10 năm 2026"); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before = Pt(18)
                p = para("Học viên"); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.runs[0].bold = True
                p = para("Quách Thành Long"); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.runs[0].italic = True
                continue
            if txt.startswith("Một phiên demo") or txt.startswith("Ngân hàng đề lớn"):
                callout("Ghi chú.", txt); continue
            para(txt)
        elif kind in ("fig", "pair", "callout", "h3"):
            render_part(part)
        elif kind == "table" and item["_orig"] is None:
            render_part(part)
        elif kind == "table":
            rep = X.TABLE_REPLACE.get(n, {}).get(tables_seen)
            tables_seen += 1
            if rep:
                table(*rep)
            else:
                rows = part[2]
                if n == 26: rows = rows + X.EXTRA_REFS
                table(part[1], rows, part[3] if len(part) > 3 else None)
    if n == 6:
        head, paras, cols, rows, note = X.SCORING
        doc.add_paragraph(renumber(head), style="Heading 2")
        for t in paras: para(t)
        table(cols, rows, [2200, 3300, 3500])
        callout("Điều kiện hoàn thành.", note)
    for f in figs:
        figure(*f)
    if n in X.STEPS:
        ttl, rows = X.STEPS[n]
        doc.add_paragraph(ttl, style="Heading 3")
        table(X.STEP_HEAD, rows, [500, 3000, 5520])
    for part in R7.APPEND.get(n, []):
        render_part(part)

# ─────────── Appendix: gallery ───────────
page_break()
chapter_no[0] = 0
para("PHỤ LỤC A", "Kicker")
doc.add_paragraph("Phụ lục A. Thư viện ảnh game", style="Heading 1")
para("Ảnh chụp tự động từ các bài kiểm thử Unity Play Mode, độ phân giải 1600×900.", "Lead")
gallery = [("shots/01_menu.png", "Menu chính"), ("shots/02_guide.png", "Cách chơi"), ("shots/g09_dialogue_choice.png", "Hội thoại N5"),
           ("shots/05_quests.png", "Nhật ký nhiệm vụ"), ("shots/g04_map.png", "Bản đồ Nihongo City"), ("shots/g08_restaurant_menu.png", "Thực đơn Sushi Hibari"),
           ("shots/g06_exam_hub.png", "Trung tâm luyện thi"), ("shots/g07_exam_play.png", "Làm bài JLPT N5"), ("shots/07_settings.png", "Cài đặt"),
           ("shots/m02_about.png", "Về tôi"), ("shots/s01_station.png", "Ga Sakura-Midori"), ("shots/s02_sushi.png", "Nhà hàng sushi"),
           ("shots/s03_classroom.png", "Lớp học Hibari"), ("shots/s04_bedroom.png", "Phòng ngủ"),
           ("shots7/city_02_konbini_front.png", "Mặt tiền ひばりマート"), ("shots7/city_06_signpost.png", "Cột chỉ đường trong phố"),
           ("shots7/bedroom_02_room_overview.png", "Phòng trọ 1K"), ("shots7/bedroom_11_morning.png", "Buổi sáng sau khi ngủ"),
           ("shots7/ui_10_station_hall.png", "Sảnh ga Hibari"), ("shots7/ui_14_platform_waiting.png", "Chờ tàu ở sân 2"),
           ("shots7/ui_19_arrived_minato.png", "Tới ga Minato"), ("shots7/ui_26_neighbour_colours.png", "Hàng xóm có màu da, tóc, áo riêng"),
           ("shots7/gamecenter_01_city_entrance.png", "Lối vào Game Center"), ("shots7/gamecenter_05_board.png", "Bàn Kana Match"),
           ("shots7/school_07_classroom.png", "Lớp học sau bài thi"), ("shots7/world_horizon_dusk_west_edge.png", "Chân trời lúc hoàng hôn")]
half = (CONTENT_W - 200) // 2
g = fixed_table((len(gallery) + 1) // 2, 2, [half + 100, half + 100]); borders(g, "FFFFFF", False); cell_margins(g, 60, 120, 60, 60)
for k, (f, cap) in enumerate(gallery):
    c = g.rows[k // 2].cells[k % 2]
    pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(str(P / _jpg(f)), width=Inches(half / 1440))
    cp = c.add_paragraph(style="Caption"); cp.paragraph_format.space_after = Pt(4)
    x = cp.add_run(f"A.{k + 1}  "); x.bold = True; x.italic = False; x.font.color.rgb = RGBColor.from_string(INK)
    cp.add_run(cap)

cp = doc.core_properties
cp.title = "NihongoLife — Báo cáo đồ án Project 3 Capstone"
cp.author = "Quách Thành Long"
cp.subject = "Game 3D mô phỏng cuộc sống kết hợp học tiếng Nhật N5"
cp.keywords = "NihongoLife, Unity, JLPT N5, VTC Academy"
doc.save(OUT)
print("saved", OUT)
