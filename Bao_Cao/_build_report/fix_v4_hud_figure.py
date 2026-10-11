# -*- coding: utf-8 -*-
"""Keeps figure numbering: removes the separate new-HUD figure and puts the new HUD image into the existing
"Hình 10.2" (left cell of the figure table under §10.1), replacing the old-HUD screenshot."""
import os
import docx
from docx.shared import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCX = os.path.join(ROOT, "NihongoLife_Bao_Cao_Capstone_v4.docx")
d = docx.Document(DOCX)

# 1. Remove the separate figure (picture paragraph + its caption) added earlier.
paras = d.paragraphs
for i, p in enumerate(paras):
    if p.text.startswith("Hình 10.2. Player HUD: portrait"):
        pic = paras[i - 1]
        assert pic._element.xpath(".//pic:pic"), "picture paragraph expected before the caption"
        pic._element.getparent().remove(pic._element)
        p._element.getparent().remove(p._element)
        print("removed separate figure")
        break

# 2. Replace the image of "Hình 10.2. HUD ở 1280×720" inside the figure table.
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            if any(par.text.startswith("Hình 10.2. HUD") for par in cell.paragraphs):
                for par in cell.paragraphs:
                    blips = par._element.xpath(".//pic:pic")
                    if not blips:
                        continue
                    extent = par._element.xpath(".//wp:extent")[0]
                    width = int(extent.get("cx"))
                    for r in list(par.runs):
                        r._element.getparent().remove(r._element)
                    par.add_run().add_picture(os.path.join(ROOT, "hud-regression", "hud_1920x1080_widget.png"), width=Emu(width))
                for par in cell.paragraphs:
                    if par.text.startswith("Hình 10.2. HUD"):
                        par.runs[0].text = "Hình 10.2. Player HUD mới: portrait từ model 3D thật, cấp độ, ví và năm vòng nhu cầu"
                        for r in par.runs[1:]:
                            r.text = ""
                print("replaced figure 10.2 image")
d.save(DOCX)
