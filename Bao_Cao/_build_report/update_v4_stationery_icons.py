# -*- coding: utf-8 -*-
"""In-place update of NihongoLife_Bao_Cao_Capstone_v4.docx for the stationery pictures rendered from real 3D models
(§6.3, §12.2 with a new Hình 12.1, Bảng 13.1). The old Hình 12.1 becomes 12.2. Idempotent."""
import copy
import os
import sys

import docx
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCX = os.path.join(ROOT, "NihongoLife_Bao_Cao_Capstone_v4.docx")

d = docx.Document(DOCX)
if any(p.text.startswith("Hình 12.1. Quầy Văn phòng phẩm") for p in d.paragraphs):
    print("stationery update already applied")
    sys.exit(0)


def find(prefix, style=None):
    for p in d.paragraphs:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)


def clone_after(anchor, template, text=None):
    el = copy.deepcopy(template._element)
    for child in list(el):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            el.remove(child)
    anchor._element.addnext(el)
    para = docx.text.paragraph.Paragraph(el, anchor._parent)
    if text is not None:
        run = para.add_run(text)
        if template.runs and template.runs[0]._element.rPr is not None:
            run._element.insert(0, copy.deepcopy(template.runs[0]._element.rPr))
    return para


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


# Templates: a normal paragraph, a figure paragraph (the one holding Hình 12.1's picture) and its caption.
old_caption = find("Hình 12.1. Quy trình một vòng phát triển", "Caption")
idx = [i for i, p in enumerate(d.paragraphs) if p._element is old_caption._element][0]
figure = d.paragraphs[idx - 1]
assert figure._element.xpath(".//w:drawing"), "figure paragraph above the caption"

# ─── 6.3: the stationery cards show real pictures ───
shop = find("Quầy Văn phòng phẩm (ぶんぼうぐ) bán giấy cho sổ tay")
set_text(shop, shop.text.rstrip() + " Thẻ hàng của gói 12 và 30 trang dùng ảnh render từ mô hình 3D sổ gáy xoắn và vở "
         "bìa cứng có sẵn trong dự án (mục 12.2); gói giấy rời 5 trang chưa có mô hình phù hợp nên tạm giữ chữ 紙.")

# ─── 12.2: stationery pictures from real models + figure ───
pipeline = find("Codex xây dựng pipeline tạo icon từ mô hình 3D")
p1 = clone_after(pipeline, pipeline,
    "Ảnh cho quầy Văn phòng phẩm dùng cùng nguyên tắc: ItemIconRenderer có thêm hai lệnh chạy batchmode, không cần thao "
    "tác menu — PreviewStationery chụp toàn bộ 41 mô hình văn phòng phẩm trong dự án để chọn bằng mắt, RenderStationery "
    "ghi ảnh sản phẩm vào Resources/Items. Mọi ảnh dùng chung góc máy 3/4, bộ đèn key/fill/rim và lề cắt 6% quanh "
    "hình bóng vật, nền trong suốt 512 px. Bộ Styloo Classroom khi nhập dạng FBX bị mất texture bảng màu nên mô hình "
    "hiện trắng; renderer đọc bảng màu ngay từ tệp GLB cùng tên trong bộ tài nguyên khi chụp, không thêm texture mới vào "
    "dự án. Gói 12 trang dùng sổ gáy xoắn PRINCIPALOFFICEnotebook_001, gói 30 trang dùng vở bìa cứng book__004 đặt nằm. "
    "Gói giấy rời 5 trang được ghi nhận là thiếu mô hình: mẫu gần nhất (postit) chỉ là một tờ giấy dán cong, không đúng "
    "sản phẩm, nên đồ án không dùng thay.")
pic = clone_after(p1, figure)
pic.add_run().add_picture(os.path.join(HERE, "img", "stationery_shelf_crop.png"), width=Cm(14))
clone_after(pic, old_caption, "Hình 12.1. Quầy Văn phòng phẩm: ảnh gói 12 và 30 trang render từ mô hình 3D thật")
set_text(old_caption, "Hình 12.2. Quy trình một vòng phát triển")

# ─── Bảng 13.1: the notebook test now checks the pictures too ───
body = list(d.element.body.iterchildren())
for k, e in enumerate(body):
    if e.tag.endswith("}p") and k + 1 < len(body) and body[k + 1].tag.endswith("}tbl") and \
            "".join(x.text or "" for x in e.iter() if x.tag.endswith("}t")).startswith("Bảng 13.1."):
        table = docx.table.Table(body[k + 1], d)
        break
for row in table.rows:
    if row.cells[0].text.startswith("HudStartupPlayModeTests"):
        cell = row.cells[1].paragraphs[0]
        if "ảnh render" not in cell.text:
            set_text(cell, cell.text.rstrip(".") + "; thẻ giấy 12/30 trang hiển thị ảnh render từ mô hình")

d.save(DOCX)
print("updated", DOCX)
