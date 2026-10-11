# -*- coding: utf-8 -*-
"""In-place update of NihongoLife_Bao_Cao_Capstone_v4.docx (keeps every manual edit): new Player HUD, notebook,
exam highlighting, stationery shelf, the HUD-at-startup fix, tests and performance. Idempotent: skips if the new
section already exists."""
import copy
import os
import sys

import docx
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCX = os.path.join(ROOT, "NihongoLife_Bao_Cao_Capstone_v4.docx")

d = docx.Document(DOCX)
P = d.paragraphs


def update_exam_table():
    """Bảng 9.1: reflect the JLPT N5 mock 2 and the local IELTS Reading packages (idempotent)."""
    doc = docx.Document(DOCX)
    body = list(doc.element.body.iterchildren())
    for k, e in enumerate(body):
        if e.tag.endswith("}p") and "".join(x.text or "" for x in e.iter() if x.tag.endswith("}t")).startswith("Bảng 9.1"):
            table = docx.table.Table(body[k + 1], doc)
            break
    else:
        return
    texts = [[c.text for c in r.cells] for r in table.rows]
    if any("đề luyện tập 2" in t[0] for t in texts):
        return

    def set_cell(cell, value):
        first = cell.paragraphs[0]
        for extra in cell.paragraphs[1:]:
            extra._element.getparent().remove(extra._element)
        if first.runs:
            first.runs[0].text = value
            for r in first.runs[1:]:
                r.text = ""
        else:
            first.add_run(value)

    rows = table.rows
    # JLPT mock 2 after mock 1.
    tr = copy.deepcopy(rows[1]._tr)
    rows[1]._tr.addnext(tr)
    new = [r for r in table.rows if r._tr is tr][0]
    set_cell(new.cells[0], "JLPT N5 — đề luyện tập 2")
    set_cell(new.cells[1], "Đề tự soạn theo cấu trúc chính thức (67 câu, theo từng もんだい); nghe đọc bằng giọng AI, ghi rõ; "
                           "chấm theo hai nhóm điểm 0–120 / 0–60, điểm liệt và ngưỡng đỗ 80/180, phân tích theo もんだい")
    for r in table.rows:
        if r.cells[0].text.startswith("IELTS Listening đầy đủ"):
            set_cell(r.cells[0], "IELTS Listening và Reading đầy đủ (gói cục bộ)")
            set_cell(r.cells[1], "Engine đọc gói đề từ LocalContent (ngoài Assets, không phát hành): 40 câu, audio hoặc bài đọc, "
                                 "Reading tính giờ 60 phút và tự nộp, hai chế độ Luyện tập/Thi, tự lưu, xem lại, tô sáng và ghi sổ tay")
        elif r.cells[0].text.startswith("IELTS Reading/Writing/Speaking"):
            set_cell(r.cells[0], "IELTS Writing/Speaking đề thật, JLPT chính thức")
    doc.save(DOCX)
    print("table 9.1 updated")



ALREADY = any(p.text.startswith("10.5. Sổ tay") for p in P)
if ALREADY:
    print("main update already applied")
    update_exam_table()
    sys.exit(0)


def find(prefix, style=None):
    for p in d.paragraphs:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)


def clone_after(anchor, template, text=None):
    """New paragraph after `anchor`, formatted like `template` (runs removed, text added in one run copied from the
    template's first run so fonts match)."""
    el = copy.deepcopy(template._element)
    for child in list(el):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            el.remove(child)
    anchor._element.addnext(el)
    para = docx.text.paragraph.Paragraph(el, anchor._parent)
    if text is not None:
        run = para.add_run(text)
        if template.runs:
            src = template.runs[0]._element.rPr
            if src is not None:
                run._element.insert(0, copy.deepcopy(src))
    return para


def picture_after(anchor, figure_template, path, width_cm):
    para = clone_after(anchor, figure_template)
    para.add_run().add_picture(path, width=Cm(width_cm))
    return para


normal = find("Player HUD mới do HUDUI")
figure = d.paragraphs[[i for i, p in enumerate(d.paragraphs) if p._element is normal._element][0] + 1]
caption = find("Hình 10.1.", "Caption")
heading2 = find("10.4. Đánh giá theo nguyên tắc khả dụng", "Heading 2")

# ─── 10.1: the new Player HUD ───
normal.runs[0].text = (
    "Player HUD do HUDUI kết nối và StatusDock trình bày, theo phong cách kính tối bo góc, gọn ở góc trái dưới: "
    "portrait đầu nhân vật được chụp từ chính model 3D đang điều khiển (HudPortrait sao chép nhánh Visual sang một "
    "sân khấu riêng ở layer 30, chụp một lần vào RenderTexture 256 px rồi tắt hẳn; chỉ chụp lại khi đổi nhân vật), "
    "badge cấp độ ở mép portrait, tên và địa điểm ở giữa, ví ¥ viền vàng bên phải, thanh EXP mảnh và năm vòng tiến "
    "trình Sức khỏe, Thể lực, No, Khát, Tỉnh táo — mỗi vòng có icon vẽ thủ tục (HudIcons) và phần trăm bên trong, "
    "không kèm nhãn chữ. Dữ liệu vẫn lấy từ PlayerStatus và PlayerInventory, giữ nguyên logic đổi màu khi chỉ số "
    "thấp và thẻ cảnh báo. Góc trái trên chỉ giữ một mục tiêu đang theo dõi; giữa phía trên dành cho zone bar và "
    "HudFeed tối đa ba thông báo; giữa phía dưới chỉ có một prompt tương tác. HUD mờ/ẩn trong hội thoại hoặc hoạt "
    "động toàn màn hình. Desktop dùng phím tắt; touch layout giữ action bar, joystick và nút tương tác.")
for r in normal.runs[1:]:
    r.text = ""
hud_fig = picture_after(caption, figure, os.path.join(ROOT, "hud-regression", "hud_1920x1080_widget.png"), 12.5)
clone_after(hud_fig, caption, "Hình 10.2. Player HUD: portrait từ model 3D thật, cấp độ, ví và năm vòng nhu cầu (ảnh chụp ở 1920×1080)")

# ─── 10.5: notebook and highlighting (end of chapter 10) ───
chapter11 = find("CHƯƠNG 11.", "Heading 1")
last = d.paragraphs[[i for i, p in enumerate(d.paragraphs) if p._element is chapter11._element][0] - 1]
h = clone_after(last, heading2, "10.5. Sổ tay và công cụ tô sáng khi làm bài")
t1 = clone_after(h, normal,
    "Sổ tay là tính năng dùng chung cho cả game (phím L, nút “Sổ tay” trên HUD hoặc trên thanh công cụ bài thi): một "
    "cuốn sổ gáy xoắn nhỏ gọn, trang giấy kẻ dòng có lề, ô tiêu đề, dãy số trang để chuyển nhanh, bộ đếm ký tự và nút "
    "tẩy trang (bấm hai lần). Số trang có giới hạn — bắt đầu với 6 trang, mỗi trang tối đa 600 ký tự, tối đa 99 trang — "
    "và nội dung được lưu cùng tiến trình (gộp bản lưu online lấy bản sửa gần nhất). Khi viết kín mọi trang, sổ chỉ "
    "đường tới quầy Văn phòng phẩm của Hibari Mart để mua thêm giấy. Sổ luôn nằm trên màn hình thi nên có thể ghi chép "
    "trong lúc làm bài.")
t2 = clone_after(t1, normal,
    "Trong bài thi IELTS và JLPT, thanh công cụ ở đầu màn hình có bút tô sáng ba màu, tẩy, “Xoá hết”, “→ Sổ” (chép các "
    "đoạn đã tô vào sổ tay) và nút mở sổ. Người làm bài kéo chuột qua chữ để tô như bài thi trên máy: văn bản tiếng Anh "
    "tự bắt trọn từ, tiếng Nhật giữ đúng các ký tự được kéo qua. Highlight được chèn bằng thẻ <mark> vào chính văn bản "
    "nên xuống dòng và cuộn đúng; khi tắt bút, chữ không bắt chuột nên việc cuộn và trả lời không bị ảnh hưởng. Bài "
    "IELTS lưu highlight cùng bài làm, mở lại vẫn còn.")
nb = picture_after(t2, figure, os.path.join(HERE, "img", "notebook_crop.png"), 7.5)
c1 = clone_after(nb, caption, "Hình 10.6. Sổ tay: trang giấy kẻ dòng, số trang và bộ đếm ký tự")
hl = picture_after(c1, figure, os.path.join(HERE, "img", "jlpt_highlight_crop.png"), 13)
clone_after(hl, caption, "Hình 10.7. Thanh công cụ tô sáng trong bài JLPT N5 (đề tự soạn), chữ 毎朝 đã được tô")

# ─── 6.3: stationery shelf ───
shop = find("Konbini ひばりマート có bốn kệ")
shop.runs[0].text = shop.runs[0].text.replace("có bốn kệ theo nhóm hàng", "có bốn kệ theo nhóm hàng và một quầy Văn phòng phẩm")
clone_after(shop, normal,
    "Quầy Văn phòng phẩm (ぶんぼうぐ) bán giấy cho sổ tay: gói 5, 12 và 30 trang giá ¥120, ¥260 và ¥580. Thanh toán xong, "
    "số trang được cộng thẳng vào sổ, không chiếm ô balo; nếu sổ đã đủ 99 trang, phần không nhận được sẽ được hoàn tiền.")

# ─── 9.2: exam tools ───
grader = find("Bộ chấm IELTS xử lý đúng quy ước")
clone_after(grader, normal,
    "Trong khi làm bài, người học có thể tô sáng đề và ghi chú vào sổ tay ngay trên màn hình thi (mục 10.5). Gói đề "
    "Cambridge chỉ dùng cục bộ cho luyện tập cá nhân, không đưa vào bản build hay báo cáo.")

# ─── 13.1 / 13.2 tables ───
def table_after_caption(prefix):
    body = list(d.element.body.iterchildren())
    for k, e in enumerate(body):
        if e.tag.endswith("}p") and "".join(x.text or "" for x in e.iter() if x.tag.endswith("}t")).startswith(prefix):
            return docx.table.Table(body[k + 1], d)
    raise KeyError(prefix)


def add_row(table, values, before_last=False):
    template = table.rows[-1]._tr
    tr = copy.deepcopy(template)
    if before_last:
        template.addprevious(tr)
    else:
        template.addnext(tr)
    row = docx.table.Table(table._tbl, d).rows[[r._tr for r in table.rows].index(tr)]
    for cell, value in zip(row.cells, values):
        paras = cell.paragraphs
        first = paras[0]
        for extra in paras[1:]:
            extra._element.getparent().remove(extra._element)
        if first.runs:
            first.runs[0].text = value
            for r in first.runs[1:]:
                r.text = ""
        else:
            first.add_run(value)


tests = table_after_caption("Bảng 13.1.")
add_row(tests, ["HudStatusWidget_RealPortrait_RingsFromStatus_Resolutions",
                "Portrait chụp từ model thật và chụp lại khi đổi nhân vật; không render lại mỗi frame; % các vòng khớp PlayerStatus; không nhãn chữ; 1080p và 768p không tràn, không đè thẻ mục tiêu",
                "Đạt"], before_last=True)
add_row(tests, ["HudStartupPlayModeTests · NotebookHighlightPlayModeTests",
                "Vào thành phố luôn thấy HUD hoặc hội thoại; sổ tay (giới hạn trang, tẩy, lưu, mua giấy); tô sáng IELTS/JLPT bằng sự kiện chuột thật, lưu theo bài làm, chép sang sổ",
                "Đạt (4 bài)"], before_last=True)
bugs = table_after_caption("Bảng 13.2.")
add_row(bugs, ["Vào thành phố không thấy HUD lẫn hội thoại",
               "Hội thoại mở đầu bắt đầu trước khi HUD đăng ký sự kiện; StatusDock ẩn vì “đang có hội thoại” nhưng hộp hội thoại chưa hiện",
               "HUDUI đăng ký lại và cho DialogueManager hiển thị lại hội thoại đang mở; thêm bài kiểm thử khởi động"])

# ─── 13.3 performance ───
perf = find("Đo trong Editor ở chế độ batch")
clone_after(perf, normal,
    "Portrait trên HUD không tốn chi phí theo frame: camera riêng chỉ render một lần cho mỗi model rồi tắt cùng sân khấu "
    "của nó. Trong bài kiểm thử, thời gian khung khi bật và khi ẩn widget chênh nhau trong phạm vi dao động đo "
    "(khoảng 1,4–2,2 ms ở cả hai trường hợp).")

# ─── 14.2 contribution ───
hud_bullet = find("– Player HUD thống nhất với portrait model thật")
clone_after(hud_bullet, hud_bullet,
    "– Sổ tay giới hạn trang gắn với kinh tế trong game (mua giấy ở konbini) và công cụ tô sáng/ghi chú ngay trong bài thi IELTS, JLPT.")

d.save(DOCX)
print("updated", DOCX)

update_exam_table()
