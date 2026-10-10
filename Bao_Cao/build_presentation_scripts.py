from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'Bao_Cao'
NAVY='17365D'; BLUE='2F5597'; PALE='EAF1F8'; LIGHT='F4F6F9'; MUTED='667085'; WHITE='FFFFFF'

def shade(cell, color):
    pr=cell._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),color); pr.append(x)
def margins(cell):
    pr=cell._tc.get_or_add_tcPr(); x=OxmlElement('w:tcMar'); pr.append(x)
    for n,v in [('top','80'),('start','110'),('bottom','80'),('end','110')]:
        z=OxmlElement('w:'+n); z.set(qn('w:w'),v); z.set(qn('w:type'),'dxa'); x.append(z)
def setup(title, subtitle):
    d=Document(); s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7); s.top_margin=Cm(2.2); s.bottom_margin=Cm(2); s.left_margin=Cm(2.5); s.right_margin=Cm(2.2)
    n=d.styles['Normal']; n.font.name='Times New Roman'; n.font.size=Pt(11); n._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); n.paragraph_format.space_after=Pt(5); n.paragraph_format.line_spacing=1.2
    for nm,sz,col in [('Title',26,NAVY),('Subtitle',13,MUTED),('Heading 1',17,NAVY),('Heading 2',13,BLUE)]:
        st=d.styles[nm]; st.font.name='Times New Roman'; st.font.size=Pt(sz); st.font.color.rgb=RGBColor.from_string(col); st.font.bold=nm!='Subtitle'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    p=d.add_paragraph(title,'Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70)
    p=d.add_paragraph(subtitle,'Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph('NihongoLife • Quách Thành Long • VTC Academy • 10/10/2026'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
    d.add_page_break(); return d
def para(d,text,bold=None):
    p=d.add_paragraph();
    if bold and text.startswith(bold): p.add_run(bold).bold=True; p.add_run(text[len(bold):])
    else:p.add_run(text)
    return p
def bullets(d,items):
    for x in items:
        p=d.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Cm(.7); p.add_run(x)
def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); margins(c); c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9)
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=str(v); cs[i].width=Cm(widths[i]); margins(cs[i]); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cs[i],LIGHT)
            for p in cs[i].paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.size=Pt(8.7)
    d.add_paragraph(); return t
def footer(d):
    p=d.sections[0].footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('NIHONGOLIFE • CAPSTONE'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

# Trailer script
d=setup('KỊCH BẢN TRAILER NIHONGOLIFE','Bản đề xuất mới • storytelling • 90 giây • 1920×1080 / 30 fps')
d.add_heading('1. Đánh giá trailer hiện tại',1)
para(d,'Trailer hiện tại dài 67,43 giây, 1920×1080, 30 fps, H.264/AAC. Cấu trúc là montage: logo → toàn cảnh Hibari → đi bộ/hội thoại → konbini → sushi → phòng trọ → ga → lớp học → Game Center/Kana Match → end card. Nhận diện rõ và nhịp gọn, nhưng chưa có một mục tiêu xuyên suốt; nhiều shot mang tính “tham quan”, chưa cho thấy Phase 5–7 như Midori Island, farming, vật nuôi, việc làm, Task Journal, level và kinh tế. Trailer mới nên kể một ngày của du học sinh, dùng Midori làm cao trào và chỉ nhắc online như khả năng mở rộng, không trình bày như tính năng đã kiểm chứng hoàn toàn.')
d.add_heading('2. Ý tưởng sáng tạo',1)
para(d,'Logline: “Một ngày mới ở Hibari-chō: học tiếng Nhật bằng cách sống, làm việc và kết nối.”')
bullets(d,['Nhịp cảm xúc: bỡ ngỡ → chủ động → thành thạo → mở rộng thế giới → thuộc về cộng đồng.','Trục hình ảnh: thành phố buổi sáng → nhiệm vụ → giao tiếp → việc làm → chuyến tàu → Midori → lớp học/kết quả → cộng đồng.','Trục âm thanh: ambience yên tĩnh, nhạc tăng dần, sound cue cho quest/Yen/thu hoạch, kết bằng chủ đề thương hiệu.','Tỷ lệ nội dung: 70% gameplay thao tác thật, 20% cinematic fly-through, 10% title/end card.'])
d.add_heading('3. Shot list 90 giây',1)
rows=[
('00:00–00:05','Cold open: ga Hibari lúc sáng; nhân vật bước xuống tàu','Wide dolly-in, 24–28 mm; cut theo tiếng chuông ga','Ambience ga, một nhịp piano','VO: “Một ngôn ngữ không chỉ nằm trong sách.”'),
('00:05–00:10','Logo NihongoLife + Hibari-chō','Logo trên nền phố; parallax nhẹ','Theme motif','Text: Học tiếng Nhật bằng cách sống.'),
('00:10–00:18','Nhận quest, bản đồ, waypoint','Over-shoulder → cận HUD mục tiêu','Quest chime','VO: “Mỗi ngày bắt đầu bằng một mục tiêu thật.”'),
('00:18–00:28','Hội thoại Tanaka; chọn sai, nhận giải thích, chọn lại','Shot/reverse-shot; cận lựa chọn; tránh giữ UI quá lâu','Voice Nhật + click mềm','VO: “Thử, sai và nói lại — không áp lực.”'),
('00:28–00:38','Konbini: bê thùng, xếp kệ, trả lời khách, nhận ¥600','Montage 4 cut; handheld rất nhẹ; match cut tay/vật phẩm','Foley thùng hàng, register ding','Text nhỏ: Việc làm nhiều bước • phần thưởng thật'),
('00:38–00:45','Sushi: đọc order và giao đúng bàn','Tracking ngang; insert món ăn/icon','Nhịp trống tăng','VO: “Ngôn ngữ trở thành hành động.”'),
('00:45–00:52','Mua vé ¥450, cửa tàu đóng, chuyển cảnh','Cận máy vé → whip pan sang tàu','Chuông cửa, rail whoosh','Text: Hibari → みどりじま'),
('00:52–01:06','Midori: toàn cảnh, xới, gieo, tưới, cây lớn, thu hoạch','Drone reveal; macro tool/crop; time-lapse 3 stage','Nhạc mở rộng, nước, harvest chime','VO: “Làm nông, học từ mới và xây dựng cuộc sống của riêng mình.”'),
('01:06–01:13','Cho bò ăn, vuốt ve; sổ tay ghi từ/achievement','Low-angle thân thiện; cận thẻ động vật','Animal + page flip','Không lời; để hình kể chuyện.'),
('01:13–01:20','Lớp học: JLPT/IELTS; kết quả và cô Morita','Push-in bàn thi → scoreboard','Nhạc giảm, success cue','VO: “Luyện tập. Đo tiến bộ. Tiếp tục.”'),
('01:20–01:26','Profile level/XP/Yen + Task Journal; trở về Hibari hoàng hôn','UI insert ngắn → wide sunset','Theme reprise','VO: “Từng việc nhỏ đưa bạn gần hơn tới cuộc sống mới.”'),
('01:26–01:30','End card/logo','Logo, tagline, thông tin Capstone','Resolve','Text: NihongoLife • ひばり町で、日本語と暮らそう')]
table(d,['Thời lượng','Hình ảnh/scene','Góc quay & dựng','Âm thanh','Lời dẫn/chữ'],rows,[2.2,4.4,3.8,2.8,4.0])
d.add_heading('4. Quy chuẩn quay và hậu kỳ',1)
bullets(d,['Khóa UI debug; dùng cùng font và palette; không quay nội dung Cambridge có bản quyền.','Cinematic shot 5–8 giây; gameplay UI 2–4 giây; ưu tiên cut theo hành động và âm thanh.','Không phóng đại online/co-op; nếu xuất hiện chỉ dùng title “hạ tầng tích hợp, đang tiếp tục kiểm chứng”.','Thu gameplay 1080p hoặc cao hơn, bitrate đủ; tránh camera xuyên vật thể, NPC T-pose, text cắt và frame có prompt thừa.','Âm lượng: VO rõ hơn nhạc khoảng 4–6 dB; duck nhạc khi có thoại Nhật; dùng phụ đề Việt ngắn.'])
footer(d); d.save(OUT/'NihongoLife_Kich_Ban_Trailer_Moi.docx')

# Gameplay defense demo
d=setup('KỊCH BẢN GAMEPLAY DEMO BẢO VỆ CAPSTONE','Run-of-show 12–15 phút • thao tác, lời thuyết minh và phương án dự phòng')
d.add_heading('1. Mục tiêu buổi demo',1)
para(d,'Chứng minh một vòng chơi hoàn chỉnh và ba luận điểm: (1) ngôn ngữ gắn với bối cảnh; (2) hệ thống progression/economy nối các hoạt động; (3) kiến trúc data-driven cho phép mở rộng từ thành phố sang Midori và khảo thí mà vẫn dùng chung save/service.')
d.add_heading('2. Chuẩn bị trước khi bảo vệ',1)
bullets(d,['Dùng một bản save demo cố định: đủ Yen mua vé/hạt, có bình tưới, inventory còn ô trống, chưa nhận ca Midori.','Mở game từ 00_Bootstrap, kiểm tra scene 60_MidoriIsland trong Build Settings, âm lượng và độ phân giải 1920×1080.','Tắt/ẩn dữ liệu Cambridge nếu máy trình chiếu không có quyền; dùng đề rút gọn tự soạn.','Chuẩn bị video dự phòng 3–5 phút cho: tàu, farming, việc làm, IELTS; không phụ thuộc Supabase/Agora trực tiếp.','Đóng ứng dụng nền và khóa thông báo; kiểm tra chuột Ctrl, ESC, N, P, F, Tab, K.'])
d.add_heading('3. Run-of-show',1)
rows=[
('00:00–00:45','Menu → chọn nhân vật → vào Hibari','“NihongoLife là life-sim giáo dục. Vòng cốt lõi không cần mạng; cloud là lớp tăng cường.”','Nếu load chậm: nói về Bootstrap/GameServices trong lúc chờ.'),
('00:45–02:00','Đi tới NPC, mở hội thoại, chọn một đáp án chưa phù hợp rồi chọn lại','“Sai không kết thúc nhiệm vụ. NPC giải thích mức lịch sự và cho cơ hội phục hồi; điểm kỹ năng ghi riêng.”','Nếu hội thoại đang suspend: nhấn R rồi tiếp tục.'),
('02:00–03:00','Mở N: Task Journal; chọn việc konbini; chỉ HUD mục tiêu','“Một sổ thống nhất cốt truyện, việc làm, nông trại, học tập và hằng ngày. Dữ liệu lấy từ progression.json.”','Không cuộn toàn bộ; chỉ mở một nhiệm vụ.'),
('03:00–04:20','Demo 2–3 bước ca konbini: bê thùng/xếp kệ/câu hỏi khách; chỉ phần thưởng ¥600','“Lương không đến từ một nút bấm. Ca làm có bước, câu hỏi, điều kiện đúng và chỉ trả một lần.”','Dùng video dự phòng nếu NPC/path bị lệch.'),
('04:20–05:10','Mở profile Tab: level/XP/Yen và nhu cầu','“XP là tổng tiến trình; level tính từ bảng. Knowledge được tách khỏi XP để không đồng nhất chơi nhiều với học giỏi.”','Đóng bằng ESC, cho thấy modal stack.'),
('05:10–06:10','Tới ga, mua vé Midori ¥450, qua cổng/lên tàu','“Giá lấy từ catalog. SceneFlow giữ save và spawn. Hệ thống có vé hỗ trợ để tránh soft-lock trên đảo.”','Nếu thời gian ngắn: dùng save ngay trước cửa tàu.'),
('06:10–08:50','Midori: mở shop P, mua hạt; xới, gieo, tưới; dùng save/cây sẵn để thu hoạch','“Ruộng có state lưu trong PlayerProgressDto.island. Cây tăng trưởng theo UTC, nên rời đảo vẫn lớn. Thao tác có thời gian và yêu cầu công cụ.”','Không chờ cây lớn trực tiếp; chuyển plot đã chuẩn bị hoặc clip time-lapse.'),
('08:50–09:40','Cho bò ăn/vuốt ve; mở notebook','“Vật nuôi dùng catalog thức ăn/từ vựng. Hiện tại là chăm sóc, chưa tuyên bố breeding hay sản phẩm.”','Chỉ demo một con vật.'),
('09:40–10:40','Bán nông sản; hoàn tất ca Midori; chỉ Yen tăng','“Đây là vòng kinh tế: làm việc hoặc thu hoạch → nhận Yen → mua vé/hạt/dụng cụ → mở hoạt động tiếp.”','Nếu job chưa nhận: chỉ bán nông sản và giải thích ca bằng ảnh.'),
('10:40–12:00','Mở Exam Center: JLPT và IELTS; chạy một câu/đoạn ngắn','“JLPT là đề luyện tự soạn. IELTS engine có practice/exam, autosave và review; nội dung có bản quyền không nằm trong repo/build.”','Không chạy full audio; không mở LocalContent nhạy cảm.'),
('12:00–13:00','Kết quả test + kiến trúc slide/ảnh','“Stage B: 19 PlayMode đạt, 1 explicit; 16 EditMode. IELTS: 6 test đạt. Midori có regression toàn vòng; số hiệu năng là Editor, không phải benchmark build.”','Mở bảng trong báo cáo nếu test runner không sẵn.'),
('13:00–14:00','Kết luận tại cảnh hoàng hôn/Hibari','“Giá trị nằm ở việc người học dùng ngôn ngữ để sống trong thế giới, nhận phản hồi và thấy tiến bộ. Phần online nhiều người và WebGL vẫn là hạng mục kiểm chứng tiếp theo.”','Kết ở end card/logo.')]
table(d,['Thời lượng','Thao tác','Lời thuyết minh chính','Fallback'],rows,[2.1,4.5,7.2,3.2])
d.add_heading('4. Câu hỏi phản biện dự kiến',1)
table(d,['Câu hỏi','Trả lời ngắn'],[
('Game chứng minh hiệu quả học tập chưa?','Chưa. Báo cáo chứng minh implementation và usability kỹ thuật; hiệu quả cần thử nghiệm người học trước/sau.'),
('Online đã chạy thật chưa?','Code và fallback có; Supabase nhiều người và Agora hai máy chưa được kiểm chứng nên không trình diễn như hoàn tất.'),
('Vì sao có cả JLPT và IELTS?','JLPT gắn trực tiếp mục tiêu tiếng Nhật; IELTS chứng minh exam engine tổng quát. Hai loại dùng repository/UI riêng và được gắn nhãn luyện tập.'),
('Kinh tế có cân bằng không?','Dòng tiền và điều kiện đã triển khai; cân bằng cuối cần telemetry/UAT. Hiện lương lần lượt ¥600/¥700/¥500 và vé Midori ¥450.'),
('Midori có phải scene runtime tạm?','Không. Scene 60_MidoriIsland được lưu thành asset đã duyệt, có collider/spawn/return flow và PlayMode regression.'),
('Dữ liệu có mất khi sync không?','Lỗi merge từng làm rơi inventory/island/quests đã sửa và có unit test; backend thật cần được kiểm chứng lại trước phát hành.')
],[5.5,11.5])
d.add_heading('5. Nguyên tắc trình bày',1)
bullets(d,['Không demo mọi tính năng. Một vòng chơi hoàn chỉnh thuyết phục hơn nhiều popup rời rạc.','Mỗi thao tác phải đi kèm một luận điểm thiết kế hoặc kỹ thuật; tránh chỉ mô tả “bấm nút này”.','Nói rõ ranh giới: đã triển khai, đã kiểm thử, chưa kiểm thử. Đây là điểm mạnh học thuật, không phải điểm yếu.','Giữ 1–2 phút dự phòng cho lỗi hoặc câu hỏi; tổng demo trực tiếp không quá 14 phút.'])
footer(d); d.save(OUT/'NihongoLife_Kich_Ban_Gameplay_Demo_Capstone.docx')
print('done')
