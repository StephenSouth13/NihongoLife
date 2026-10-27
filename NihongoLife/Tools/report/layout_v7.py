"""Layout pass on the built v7 report: no blank paragraph before a forced page break, list fields end on their last
line, and the chapters named on the command line flow on from the previous page instead of opening a new one."""
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

path = sys.argv[1]
flow = sys.argv[2:]
d = docx.Document(path)
body = d.element.body


def text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()


removed = 0
for p in list(body.iter(qn('w:p'))):
    ppr = p.find(qn('w:pPr'))
    if ppr is None or ppr.find(qn('w:pageBreakBefore')) is None:
        continue
    prev = p.getprevious()
    while prev is not None and prev.tag == qn('w:p') and not text(prev) and prev.find('.//' + qn('w:drawing')) is None \
            and prev.find('.//' + qn('w:fldChar')) is None and prev.find('.//' + qn('w:sectPr')) is None:
        nxt = prev.getprevious()
        if prev.getparent() is body:
            body.remove(prev)
            removed += 1
        prev = nxt

# A TOC / list ends with a paragraph holding only the field's end mark; when the list fills its page that empty
# paragraph starts a page of its own. Move the mark into the list's last line.
merged = 0
for p in list(body.iter(qn('w:p'))):
    runs = p.findall(qn('w:r'))
    if len(runs) != 1 or text(p):
        continue
    fc = runs[0].find(qn('w:fldChar'))
    prev = p.getprevious()
    if fc is None or fc.get(qn('w:fldCharType')) != 'end' or prev is None or prev.tag != qn('w:p'):
        continue
    st = prev.find('.//' + qn('w:pStyle'))
    if st is None or not st.get(qn('w:val')).upper().startswith('TOC'):
        continue
    prev.append(runs[0])
    p.getparent().remove(p)
    merged += 1

changed = 0
for par in d.paragraphs:
    if par.style.name == 'Heading 1' and any(par.text.strip().startswith(f) for f in flow):
        ppr = par._p.get_or_add_pPr()
        pb = ppr.find(qn('w:pageBreakBefore'))
        if pb is not None:
            ppr.remove(pb)
        par.paragraph_format.space_before = Pt(30)
        changed += 1
# The list of tables follows the list of figures on the same page (both are short lists of one line each).
for par in d.paragraphs:
    if par.text.strip() == 'DANH MỤC BẢNG':
        ppr = par._p.get_or_add_pPr()
        pb = ppr.find(qn('w:pageBreakBefore'))
        if pb is not None:
            ppr.remove(pb)
            par.paragraph_format.space_before = Pt(24)
            changed += 1
d.save(path)
print('blank paragraphs removed', removed, '· field ends merged', merged, '· chapters flowing', changed)
