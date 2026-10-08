import os
import sys

def create_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("openpyxl module not found. Please install it using: pip install openpyxl")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NihongoLife Timeline"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY_HEADER = "1F4E78"
    WHITE_TEXT = "FFFFFF"
    TITLE_BG = "002060"
    SUBHEADER_BG = "D9E1F2"
    SECTION_BG = "8EA9DB"
    
    # Priority Fills
    P0_FILL = "FADBD8" # Soft Red
    P1_FILL = "FCF3CF" # Soft Yellow
    P2_FILL = "D6EAF8" # Soft Blue
    P3_FILL = "EAECEE" # Soft Gray

    # Status Fills
    DONE_FILL = "D4EFDF" # Soft Green
    IN_PROGRESS_FILL = "FCF3CF" # Yellow
    TODO_FILL = "F2F4F4" # Light Gray

    font_title = Font(name="Times New Roman", size=16, bold=True, color=WHITE_TEXT)
    font_subtitle = Font(name="Times New Roman", size=11, italic=True, color="D9D9D9")
    font_header = Font(name="Times New Roman", size=11, bold=True, color=WHITE_TEXT)
    font_section = Font(name="Times New Roman", size=12, bold=True, color="002060")
    font_kpi_num = Font(name="Times New Roman", size=18, bold=True, color="1F4E78")
    font_kpi_label = Font(name="Times New Roman", size=10, bold=True, color="595959")
    font_regular = Font(name="Times New Roman", size=11)
    font_bold = Font(name="Times New Roman", size=11, bold=True)

    thin_border_side = Side(style='thin', color='D9D9D9')
    thick_bottom_side = Side(style='medium', color='1F4E78')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)

    # 1. Title Banner
    ws.merge_cells("A1:K1")
    ws["A1"] = "DỰ ÁN NIHONGOLIFE - BẢNG KHẾ HOẠCH TIMELINE & DEADLINE CHI TIẾT"
    ws["A1"].font = font_title
    ws["A1"].fill = PatternFill(start_color=TITLE_BG, end_color=TITLE_BG, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    ws.merge_cells("A2:K2")
    ws["A2"] = "Người thực hiện Duy nhất: Quách Thành Long  |  Đơn vị: VTC Academy  |  Cập nhật: 08/10/2026"
    ws["A2"].font = font_subtitle
    ws["A2"].fill = PatternFill(start_color=TITLE_BG, end_color=TITLE_BG, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # 2. KPI Summary Cards
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 30

    kpis = [
        ("A4", "A5", "B4", "B5", "TỔNG SỐ TASK", "=COUNTA(C8:C65)", "1F4E78"),
        ("D4", "D5", "E4", "E5", "ĐÃ HOÀN THÀNH", '=COUNTIF(F8:F65, "Hoàn thành")', "27AE60"),
        ("G4", "G5", "H4", "H5", "ĐANG THỰC HIỆN", '=COUNTIF(F8:F65, "Đang thực hiện")', "D4AC0D"),
        ("J4", "J5", "K4", "K5", "TỔNG GIỜ CÔNG", "=SUM(G8:G65)", "8E44AD")
    ]

    for top_l, bot_l, top_r, bot_r, label, formula, col_hex in kpis:
        ws.merge_cells(f"{top_l}:{top_r}")
        ws.merge_cells(f"{bot_l}:{bot_r}")
        ws[top_l] = label
        ws[top_l].font = font_kpi_label
        ws[top_l].alignment = Alignment(horizontal="center", vertical="center")
        ws[top_l].fill = PatternFill(start_color=SUBHEADER_BG, end_color=SUBHEADER_BG, fill_type="solid")
        
        ws[bot_l] = formula
        ws[bot_l].font = Font(name="Times New Roman", size=16, bold=True, color=col_hex)
        ws[bot_l].alignment = Alignment(horizontal="center", vertical="center")
        ws[bot_l].fill = PatternFill(start_color="F2F4F4", end_color="F2F4F4", fill_type="solid")
        
        for r in range(4, 6):
            for col in range(openpyxl.utils.column_index_from_string(top_l[0]), openpyxl.utils.column_index_from_string(top_r[0]) + 1):
                ws.cell(row=r, column=col).border = cell_border

    # 3. Table Headers
    headers = [
        ("STT", 6),
        ("Hạng mục / Module", 22),
        ("Tên Task / Chi tiết công việc", 45),
        ("Ưu tiên", 12),
        ("Hộp kiểm", 10),
        ("Trạng thái", 16),
        ("Thời gian (Giờ)", 15),
        ("Target Deadline", 16),
        ("Timeline (Tuần)", 16),
        ("Phân công", 18),
        ("Ghi chú & Nguồn yêu cầu", 42)
    ]

    ws.row_dimensions[7].height = 28
    for col_idx, (h_text, width) in enumerate(headers, 1):
        cell = ws.cell(row=7, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Data Source Tasks
    raw_tasks = [
        # P0 - Core & Foundation
        ("Core & System", "Kiểm tra & xác nhận InputSystem_Actions binding ổn định trên Windows/WebGL", "P0", "Đang thực hiện", "FALSE", 8, "12/10/2026", "Tuần 1 (T10)", "Quách Thành Long", "Đảm bảo di chuyển, camera và tương tác chạy mượt trên WebGL"),
        ("Core & System", "Chạy thử thủ công từ 00_Bootstrap đến 90_TestSandbox toàn bộ flow", "P0", "Đang thực hiện", "FALSE", 12, "14/10/2026", "Tuần 1 (T10)", "Quách Thành Long", "Test di chuyển, tương tác, dialogue, choice & Result UI"),
        ("Core & System", "Bổ sung xử lý lỗi khi asset visual bắt buộc không tồn tại (Production fallback)", "P0", "Chưa bắt đầu", "FALSE", 6, "16/10/2026", "Tuần 1 (T10)", "Quách Thành Long", "Không âm thầm thay bằng primitive trong bản chính thức"),
        ("Core & System", "Kiểm tra collider, layer, trigger và điểm spawn NPC/Item trong scene", "P0", "Chưa bắt đầu", "FALSE", 10, "18/10/2026", "Tuần 2 (T10)", "Quách Thành Long", "Đảm bảo người chơi luôn tiếp cận đúng mục tiêu không bị kẹt"),
        ("Core & System", "Bake NavMesh / AI Navigation cho các khu vực NPC di chuyển", "P0", "Chưa bắt đầu", "FALSE", 8, "20/10/2026", "Tuần 2 (T10)", "Quách Thành Long", "Loại bỏ warning navmesh khi NPC thực hiện pathfinding"),

        # P0 - Pháp lý & Định hướng
        ("Legal & Business", "Xác minh điều khoản thương mại (License) của asset Sushi Restaurant & StylooClassroom", "P0", "Chưa bắt đầu", "FALSE", 4, "22/10/2026", "Tuần 2 (T10)", "Quách Thành Long", "Rà soát EULA thương mại của các pack asset mua ngoài"),
        ("Legal & Business", "Đánh giá rủi ro pháp lý thương hiệu khi dùng nhãn hiệu 'JLPT' và 'IELTS'", "P0", "Chưa bắt đầu", "FALSE", 6, "24/10/2026", "Tuần 2 (T10)", "Quách Thành Long", "Cần có disclaimer hoặc chuyển tên chuẩn hóa theo quy định"),
        ("Legal & Business", "Gộp toàn bộ hệ thống thi thành 1 luồng duy nhất (ExamManager)", "P0", "Hoàn thành", "TRUE", 8, "08/10/2026", "Tuần 1 (T10)", "Quách Thành Long", "Đã xóa AssessmentEngine trùng lặp, giữ 1 pipeline duy nhất"),
        ("Legal & Business", "Xây dựng mô hình chi phí API (Gemini / AI Spoken) trước khi lên kế hoạch Freemium", "P0", "Chưa bắt đầu", "FALSE", 8, "26/10/2026", "Tuần 3 (T10)", "Quách Thành Long", "Tính toán token cost per user per session"),

        # P1 - Gameplay & Visual Experience
        ("Gameplay & Graphics", "Thay NPC capsule màu magenta bằng Character Prefab low-poly thật", "P1", "Chưa bắt đầu", "FALSE", 16, "28/10/2026", "Tuần 3 (T10)", "Quách Thành Long", "Dùng asset nhân vật Nhật Bản phong cách low-poly"),
        ("Gameplay & Graphics", "Gắn Mecanim Animations cho NPC: Idle, Walk, Talk, Checkout Gesture", "P1", "Chưa bắt đầu", "FALSE", 20, "31/10/2026", "Tuần 3 (T10)", "Quách Thành Long", "Tạo Animator Controller mượt mà khi hội thoại"),
        ("Gameplay & Graphics", "Tách Gameplay Collider khỏi Visual Prefab bằng Layer/Prefab rõ ràng", "P1", "Chưa bắt đầu", "FALSE", 8, "02/11/2026", "Tuần 4 (T11)", "Quách Thành Long", "Tối ưu hóa collision check và physics update"),
        ("Gameplay & Graphics", "Hiển thị Objective active/completed nhất quán khi chuyển Scene & kết thúc", "P1", "Chưa bắt đầu", "FALSE", 8, "04/11/2026", "Tuần 4 (T11)", "Quách Thành Long", "Sửa lỗi UI Objective bị trôi hoặc lệch state"),
        ("Gameplay & Graphics", "Hoàn thiện Pause Menu, Restart Scenario và Quay lại Main Menu", "P1", "Chưa bắt đầu", "FALSE", 10, "06/11/2026", "Tuần 4 (T11)", "Quách Thành Long", "Bảo đảm hủy event subscription đúng cách không gây memory leak"),
        ("Gameplay & Graphics", "Thêm visual & audio feedback cho sai vật phẩm, sai lựa chọn & sai vùng đích", "P1", "Chưa bắt đầu", "FALSE", 10, "08/11/2026", "Tuần 4 (T11)", "Quách Thành Long", "Hiệu ứng rung nhẹ screen hoặc âm thanh báo hiệu sai"),
        ("Gameplay & Graphics", "Tối ưu Camera Collision, Cursor Lock & trải nghiệm điều khiển Phím/Chuột", "P1", "Chưa bắt đầu", "FALSE", 8, "10/11/2026", "Tuần 5 (T11)", "Quách Thành Long", "Camera không bị xuyên tường khi ở góc hẹp"),

        # P1 - Audio System
        ("Audio System", "Tích hợp Audio Service phát lại Voice Clips hội thoại tiếng Nhật native", "P1", "Chưa bắt đầu", "FALSE", 16, "12/11/2026", "Tuần 5 (T11)", "Quách Thành Long", "Liên kết voiceClip với Dialogue Node graph"),
        ("Audio System", "Thêm nhạc nền (BGM), Ambient sound từng địa điểm (Konbini, Ga tàu, Quán Ramen)", "P1", "Chưa bắt đầu", "FALSE", 12, "14/11/2026", "Tuần 5 (T11)", "Quách Thành Long", "Có hiệu ứng Crossfade khi chuyển đổi địa điểm"),
        ("Audio System", "Hoàn thiện Volume Settings (BGM, Master, SFX, Voice) trong UI Menu", "P1", "Chưa bắt đầu", "FALSE", 6, "16/11/2026", "Tuần 5 (T11)", "Quách Thành Long", "Lưu cài đặt vào PlayerPrefs / Profile Local"),

        # P1 - Content & Scenario System
        ("Content & Scenario", "Rà soát toàn bộ tiếng Nhật, Reading, Romaji và bản dịch tiếng Việt", "P1", "Chưa bắt đầu", "FALSE", 16, "18/11/2026", "Tuần 6 (T11)", "Quách Thành Long", "Đảm bảo tính chính xác về mặt ngữ pháp & văn hóa Nhật"),
        ("Content & Scenario", "Playtest runtime Scenario Chapter 3: Gọi món Ramen ở nhà hàng (Order Ramen)", "P1", "Đang thực hiện", "FALSE", 12, "20/11/2026", "Tuần 6 (T11)", "Quách Thành Long", "Draft asset đã xong (`scenario_restaurant_order_ramen.asset`), cần test flow"),
        ("Content & Scenario", "Playtest runtime Scenario Chapter 4: Mua vé & hỏi đường ở nhà ga (Train Station)", "P1", "Đang thực hiện", "FALSE", 12, "22/11/2026", "Tuần 6 (T11)", "Quách Thành Long", "Draft asset đã xong (`scenario_station_buy_ticket.asset`), cần test flow"),
        ("Content & Scenario", "Playtest runtime Scenario Lễ hội mùa hè (Town Summer Festival)", "P1", "Đang thực hiện", "FALSE", 14, "24/11/2026", "Tuần 6 (T11)", "Quách Thành Long", "Draft asset đã xong (`scenario_town_summer_festival.asset`), quy tụ toàn bộ NPC"),
        ("Content & Scenario", "Tạo Checklist QA Nội dung cho mỗi scenario mới ra mắt", "P1", "Chưa bắt đầu", "FALSE", 6, "26/11/2026", "Tuần 7 (T11)", "Quách Thành Long", "Biểu mẫu kiểm thử nội dung trước khi publish"),

        # P2 - Data, Cloud & Web Integration
        ("Web & Data API", "Thử nghiệm & kiểm tra Local Save trên Windows và WebGL (IndexedDB)", "P2", "Chưa bắt đầu", "FALSE", 10, "28/11/2026", "Tuần 7 (T11)", "Quách Thành Long", "Đảm bảo lưu tiến độ học không bị mất khi F5 trang WebGL"),
        ("Web & Data API", "Thiết kế API Contract cho ScoreBreakdownDto, Progress và Session Launch Token", "P2", "Chưa bắt đầu", "FALSE", 12, "30/11/2026", "Tuần 7 (T11)", "Quách Thành Long", "Định nghĩa JSON DTO chuẩn giữa Unity & Next.js"),
        ("Web & Data API", "Thay LocalProgressRepository bằng Remote Repository gửi HTTPS REST requests", "P2", "Chưa bắt đầu", "FALSE", 16, "02/12/2026", "Tuần 8 (T12)", "Quách Thành Long", "Kết nối Unity WebGL với Backend Supabase"),
        ("Web & Data API", "Kết nối Next.js Portal với WebGL Frame thông qua Session Token", "P2", "Chưa bắt đầu", "FALSE", 16, "05/12/2026", "Tuần 8 (T12)", "Quách Thành Long", "Mở game từ Web Dashboard -> Truyền token -> Tải Scenario"),
        ("Web & Data API", "Bổ sung xác thực Token & Chống gửi điểm giả (Anti-cheat Score Validation)", "P2", "Chưa bắt đầu", "FALSE", 12, "08/12/2026", "Tuần 8 (T12)", "Quách Thành Long", "Xác nhận giao dịch an toàn phía Supabase Server Edge Function"),

        # P2 - Build, Performance & CI/CD
        ("Build & Deployment", "Thiết lập Build Profile Windows & WebGL Reproducible", "P2", "Chưa bắt đầu", "FALSE", 10, "10/12/2026", "Tuần 9 (T12)", "Quách Thành Long", "Tự động hóa build WebGL bundle chuẩn nén Brotli/Gzip"),
        ("Build & Deployment", "Tối ưu dung lượng Texture/Model/Audio & Loading Screen mượt mà", "P2", "Chưa bắt đầu", "FALSE", 12, "12/12/2026", "Tuần 9 (T12)", "Quách Thành Long", "Giảm dung lượng WebGL build < 50MB để tải nhanh"),
        ("Build & Deployment", "Kiểm tra & vá Memory Leak từ Events, Choice Buttons & Scene Reload", "P2", "Chưa bắt đầu", "FALSE", 10, "14/12/2026", "Tuần 9 (T12)", "Quách Thành Long", "Chạy Profiler phát hiện rò rỉ bộ nhớ khi chơi lâu"),
        ("Build & Deployment", "Thiết lập CI/CD chạy tự động Compile, Tests & Scenario Validation", "P2", "Chưa bắt đầu", "FALSE", 14, "16/12/2026", "Tuần 9 (T12)", "Quách Thành Long", "GitHub Actions tự động chạy Unity EditMode/PlayMode tests"),

        # P3 - Tooling & Authoring
        ("Tools & Authoring", "Tạo Editor Validator hiển thị lỗi trực tiếp trên Inspector ScenarioDefinition", "P3", "Chưa bắt đầu", "FALSE", 12, "18/12/2026", "Tuần 10 (T12)", "Quách Thành Long", "Cảnh báo ngay node ID trùng hoặc bị đứt nhánh"),
        ("Tools & Authoring", "Tạo công cụ Preview Dialogue Graph trong Unity Editor Window", "P3", "Chưa bắt đầu", "FALSE", 16, "20/12/2026", "Tuần 10 (T12)", "Quách Thành Long", "Trực quan hóa cây hội thoại cho Content Designer"),
        ("Tools & Authoring", "Viết hướng dẫn đóng góp Content & Quy trình Review trong Docs/CONTENT_GUIDE.md", "P3", "Hoàn thành", "TRUE", 8, "08/10/2026", "Tuần 1 (T10)", "Quách Thành Long", "Đã viết tài liệu Content Guide chi tiết")
    ]

    current_row = 8
    for idx, task in enumerate(raw_tasks, 1):
        mod, title, prio, status, chk, hrs, deadline, week, dev, note = task
        
        ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=2, value=mod).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=3, value=title).alignment = Alignment(horizontal="left", vertical="center")
        
        prio_cell = ws.cell(row=current_row, column=4, value=prio)
        prio_cell.alignment = Alignment(horizontal="center", vertical="center")
        if prio == "P0":
            prio_cell.fill = PatternFill(start_color=P0_FILL, end_color=P0_FILL, fill_type="solid")
            prio_cell.font = Font(name="Times New Roman", size=11, bold=True, color="900C3F")
        elif prio == "P1":
            prio_cell.fill = PatternFill(start_color=P1_FILL, end_color=P1_FILL, fill_type="solid")
            prio_cell.font = Font(name="Times New Roman", size=11, bold=True, color="7D6608")
        elif prio == "P2":
            prio_cell.fill = PatternFill(start_color=P2_FILL, end_color=P2_FILL, fill_type="solid")
            prio_cell.font = Font(name="Times New Roman", size=11, bold=True, color="1B4F72")
        else:
            prio_cell.fill = PatternFill(start_color=P3_FILL, end_color=P3_FILL, fill_type="solid")
            prio_cell.font = Font(name="Times New Roman", size=11, color="515A5A")

        chk_cell = ws.cell(row=current_row, column=5, value="☑" if chk == "TRUE" else "☐")
        chk_cell.alignment = Alignment(horizontal="center", vertical="center")
        chk_cell.font = Font(name="Times New Roman", size=12, bold=True)

        status_cell = ws.cell(row=current_row, column=6, value=status)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        if status == "Hoàn thành":
            status_cell.fill = PatternFill(start_color=DONE_FILL, end_color=DONE_FILL, fill_type="solid")
            status_cell.font = Font(name="Times New Roman", size=11, bold=True, color="1E8449")
        elif status == "Đang thực hiện":
            status_cell.fill = PatternFill(start_color=IN_PROGRESS_FILL, end_color=IN_PROGRESS_FILL, fill_type="solid")
            status_cell.font = Font(name="Times New Roman", size=11, bold=True, color="B7950B")
        else:
            status_cell.fill = PatternFill(start_color=TODO_FILL, end_color=TODO_FILL, fill_type="solid")
            status_cell.font = Font(name="Times New Roman", size=11, color="5D6D7E")

        ws.cell(row=current_row, column=7, value=hrs).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=8, value=deadline).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=9, value=week).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=10, value=dev).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=11, value=note).alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[current_row].height = 24

        for col in range(1, 12):
            c = ws.cell(row=current_row, column=col)
            c.border = cell_border
            if col not in [4, 5, 6]:
                c.font = font_regular

        current_row += 1

    # Data Validation for Status
    dv = DataValidation(type="list", formula1='"Chưa bắt đầu,Đang thực hiện,Hoàn thành"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F8:F{current_row-1}")

    output_path = "NihongoLife_Timeline_QuachThanhLong.xlsx"
    wb.save(output_path)
    print(f"File Excel timeline successfully created: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_excel()
