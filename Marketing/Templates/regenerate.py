#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Dựng lại ảnh Facebook, PDF kế hoạch và file lịch từ mẫu HTML.

Chạy từ gốc repo:
    python3 Marketing/Templates/regenerate.py
    python3 Marketing/Templates/regenerate.py --only-pdf
    python3 Marketing/Templates/regenerate.py --only-images
"""
from __future__ import annotations

import argparse
import html
import shutil
import subprocess
import tempfile
import time
from datetime import date, timedelta
from pathlib import Path

import markdown
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from chi_tiet import WHEN, render as render_chi_tiet
from nam import annual_rows, write_tree
from posts import POSTS, WEEKS
from tukhoa import render_markdown as render_tukhoa

ROOT = Path(__file__).resolve().parents[2]
MKT = ROOT / "Marketing"
TPL = MKT / "Templates"
CARD = (TPL / "card.html").read_text(encoding="utf-8")
DOW = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"]
START = date(2026, 10, 26)


def chrome(args: list[str], wait_for: Path, timeout: float, min_size: int) -> None:
    ud = Path(tempfile.mkdtemp(prefix="nlchrome-"))
    cmd = [
        "google-chrome", "--headless=new", "--no-sandbox", "--disable-gpu",
        "--disable-dev-shm-usage", "--hide-scrollbars", "--force-device-scale-factor=1",
        f"--user-data-dir={ud}", *args,
    ]
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    deadline = time.time() + timeout
    last = -1
    stable = 0
    try:
        while time.time() < deadline:
            if wait_for.exists():
                size = wait_for.stat().st_size
                if size >= min_size and size == last:
                    stable += 1
                    if stable >= 2:
                        break
                else:
                    stable = 0
                last = size
            time.sleep(0.35)
    finally:
        proc.kill()
        try:
            proc.wait(timeout=3)
        except subprocess.TimeoutExpired:
            pass
        shutil.rmtree(ud, ignore_errors=True)
    if not wait_for.exists() or wait_for.stat().st_size < min_size:
        raise RuntimeError(f"Chrome không xuất được {wait_for}")


def render_png(html_path: Path, png_path: Path, width: int, height: int) -> None:
    png_path.parent.mkdir(parents=True, exist_ok=True)
    if png_path.exists():
        png_path.unlink()
    chrome(
        ["--window-size=%d,%d" % (width, height), f"--screenshot={png_path}", html_path.as_uri()],
        png_path, timeout=30, min_size=8000,
    )
    with Image.open(png_path) as im:
        if im.size != (width, height):
            im.convert("RGB").resize((width, height), Image.Resampling.LANCZOS).save(png_path)


def quantize(path: Path) -> None:
    out = path.with_suffix(".quant.png")
    proc = subprocess.run(
        ["pngquant", "--quality=68-92", "--speed", "1", "--force", "-o", str(out), str(path)],
        capture_output=True,
    )
    if proc.returncode == 0 and out.exists() and out.stat().st_size < path.stat().st_size:
        out.replace(path)
    elif out.exists():
        out.unlink()


def shot_uri(post: dict, build: Path) -> str:
    if post.get("layout") in ("type", "need") or not post.get("shot"):
        return "data:image/gif;base64,R0lGODlhAQABAAAAACwAAAAAAQABAAA="
    src = ROOT / post["shot"]
    trim = post.get("trim")
    if not trim:
        return src.resolve().as_uri()
    top, right, bottom, left = trim
    with Image.open(src) as im:
        im = im.convert("RGB")
        w, h = im.size
        cropped = im.crop((int(w * left), int(h * top), int(w * (1 - right)), int(h * (1 - bottom))))
        out = build / f"{post['week']:02d}-{post['slug']}-shot.jpg"
        cropped.save(out, quality=86, optimize=True)
    return out.resolve().as_uri()


def card_html(post: dict, kind: str, shot: str) -> str:
    lines = "".join(f'<span class="ln">{html.escape(line)}</span>' for line in post["lines"])
    bullets = ""
    if post.get("bullets"):
        bullets = '<div class="bullets">' + "".join(
            f"<span>{html.escape(item)}</span>" for item in post["bullets"]
        ) + "</div>"
    size = "size-s" if max(len(line) for line in post["lines"]) > 20 else ""
    cls = {"sq": "sq", "pt": "pt", "st": "st"}[kind]
    if post.get("layout") in ("type", "need"):
        cls += " type"
    if post.get("layout") == "need":
        cls += " need"
    return (CARD
            .replace("{{TITLE}}", html.escape(post["hook"]))
            .replace("{{CLASS}}", cls)
            .replace("{{SIZE}}", size)
            .replace("{{POS}}", post.get("pos", "center 45%"))
            .replace("{{ZOOM}}", str(post.get("zoom", "1.06")))
            .replace("{{SHOT}}", shot)
            .replace("{{WEEK}}", f"{post['week']:02d}")
            .replace("{{KICKER}}", html.escape(post["kicker"]))
            .replace("{{LINES}}", lines)
            .replace("{{SUB}}", html.escape(post["sub"]))
            .replace("{{BULLETS}}", bullets)
            .replace("{{CTA}}", html.escape(post["cta"]))
            .replace("{{META}}", html.escape(post["meta"])))


def write_posts(build: Path, images: bool, slugs: list[str] | None = None) -> None:
    for post in POSTS:
        folder = MKT / f"Tuan-{post['week']:02d}" / post["slug"]
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "caption.txt").write_text(post["caption"].strip() + "\n", encoding="utf-8")
        (folder / "chi-tiet.md").write_text(render_chi_tiet(post, START), encoding="utf-8")
        can_quay = folder / "CAN_QUAY.md"
        if post.get("footage") or post.get("layout") == "need":
            rows_src = list(post.get("footage") or [])
            if post.get("layout") == "need":
                rows_src.insert(0, (
                    "8 giây", "Nhân vật bé đang chọn", "Haruto là mặc định, hoặc bé đang chọn",
                    "CẦN ẢNH MỚI", "Không dùng lại ảnh người lớn Tanaka/Remy.",
                ))
            rows = "\n".join(
                f"| {html.escape(a)} | {html.escape(b)} | {html.escape(c)} | {html.escape(d)} | {html.escape(e)} |"
                for a, b, c, d, e in rows_src
            )
            shot = post.get("shot") or "(không có — thẻ chữ, không lấy cảnh khác thế vào)"
            old = post.get("shot_source") or ""
            dropped = f"\nẢnh cũ đã bỏ: `{old}`. Trong đó là nhân vật người lớn trước khi đổi sang tám bé.\n" if post.get("layout") == "need" and old and not old.startswith("(") else ""
            can_quay.write_text(
                f"""# Cần quay

Ảnh đăng Facebook của bài này {'là thẻ chữ ghi CẦN ẢNH MỚI, vì chưa có ảnh nhân vật bé hoặc cảnh này' if post.get('layout') == 'need' else 'dùng ảnh chụp thật: `' + shot + '`'}.{dropped}

Video chưa có trong git. Quay bằng `TrailerReel` / test `Trailer_RecordFrames`. File dựng của bài để `Marketing/Video/Theo-Tuan/Tuan-{post['week']:02d}/`. Khung hình thô và mp4 gốc để `D:/VTC_Academy/NihongoLife/_work/marketing/`. Không commit file nặng.

| Thời lượng | Cảnh | Góc máy | Phụ đề | Ghi chú |
|---|---|---|---|---|
{rows}

Không chèn UI giả, không thêm điểm, không thêm người chơi không có trên màn hình.
""",
                encoding="utf-8",
            )
        elif can_quay.exists():
            can_quay.unlink()
        if not images or (slugs and post["slug"] not in slugs):
            continue
        jobs = [("sq", "fb-1080.png", 1080, 1080), ("pt", "fb-1350.png", 1080, 1350)]
        if post.get("story"):
            jobs.append(("st", "story-1920.png", 1080, 1920))
        for kind, name, w, h in jobs:
            html_path = build / f"{post['week']:02d}-{post['slug']}-{kind}.html"
            page = card_html(post, kind, shot_uri(post, build))
            # HTML nằm trong thư mục tạm. CSS phải trỏ tuyệt đối, nếu không Chrome vẽ trang trắng và hiện cả dòng CẦN ẢNH MỚI.
            page = page.replace('href="cards.css"', f'href="{(TPL / "cards.css").resolve().as_uri()}"')
            html_path.write_text(page, encoding="utf-8")
            dest = folder / name
            render_png(html_path, dest, w, h)
            quantize(dest)
            print(name, dest.stat().st_size, flush=True)


def write_week_readmes() -> None:
    by_week: dict[int, list] = {w["n"]: [] for w in WEEKS}
    for post in POSTS:
        by_week[post["week"]].append(post)
    for week in WEEKS:
        n = week["n"]
        lines = [
            f"# Tuần {n:02d} — {week['theme']}",
            "",
            f"Giai đoạn: {week['phase']}.",
            "",
            f"Mục tiêu: {week['goal']}",
            "",
            f"Mốc tuần (là mục tiêu, không phải số đã đạt): {week['kpi']}",
            "",
            "## Việc của từng người",
            "",
            f"- Long: {week['long']}",
            f"- Huy: {week['huy']}",
            f"- Dũng: {week['dung']}",
            "",
            "## Bài đăng",
            "",
            "Mở thư mục bài, đọc `chi-tiet.md` (giờ đăng, ba hook, cách trả lời, cách đo), rồi copy `caption.txt`. Ảnh `fb-1080.png` hoặc `fb-1350.png`. Story dùng `story-1920.png` nếu có. Thay `{{LINK_CHOI}}`, `{{LINK_FANPAGE}}`, `{{LINK_DISCORD}}` trước khi dán.",
            "",
            f"Clip tuần, nếu quay, để ở `Video/Theo-Tuan/Tuan-{n:02d}/`. Mp4 không commit. Bài nào cần clip và quy cách nằm trong README thư mục đó.",
            "",
        ]
        for post in by_week[n]:
            files = ["chi-tiet.md", "caption.txt", "fb-1080.png", "fb-1350.png"]
            if post.get("story"):
                files.append("story-1920.png")
            if post.get("footage"):
                files.append("CAN_QUAY.md")
            when = START + timedelta(days=(n - 1) * 7 + post["dow"])
            lines.append(f"### {when.strftime('%d/%m')} · {post['slug']}")
            lines.append("")
            lines.append(f"- Hook: {post['hook']}")
            lines.append(f"- File: {', '.join(files)}")
            lines.append("")
        folder = MKT / f"Tuan-{n:02d}"
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "README.md").write_text("\n".join(lines), encoding="utf-8")


def calendar_rows() -> list[dict]:
    posts_by = {(p["week"], p["dow"]): p for p in POSTS}
    rows = []
    for week in WEEKS:
        n = week["n"]
        for dow in range(7):
            day = START + timedelta(days=(n - 1) * 7 + dow)
            post = posts_by.get((n, dow))
            if post:
                folder = f"Tuan-{n:02d}/{post['slug']}"
                asset = f"{folder}/fb-1080.png"
                rows.append({
                    "date": day, "time": WHEN[post["dow"]][0], "week": n, "channel": "Facebook",
                    "kind": "Ảnh 1:1 và 4:5" + (" + story 9:16" if post.get("story") else ""),
                    "topic": post["kicker"], "hook": post["hook"],
                    "caption": f"{folder}/caption.txt", "asset": asset,
                    "cta": post["cta"], "owner": "Huy",
                    "kpi": week["kpi"],
                    "status": "Cần quay" if post.get("layout") == "need" else "Sẵn sàng — điền link",
                })
            if dow == 0:
                rows.append({
                    "date": day, "time": "09:00", "week": n, "channel": "Discord",
                    "kind": "Tin chữ", "topic": week["theme"],
                    "hook": f"Tuần này: {week['theme']}. {week['goal']}",
                    "caption": f"Tuan-{n:02d}/README.md", "asset": "—",
                    "cta": "Vào giờ chơi nếu là tuần có hẹn", "owner": "Huy",
                    "kpi": week["kpi"], "status": "Việc nội bộ",
                })
            if dow == 2 and n < 3:
                rows.append({
                    "date": day, "time": "12:30", "week": n, "channel": "Facebook",
                    "kind": "Tương tác, không bài mới", "topic": "Trả lời bình luận",
                    "hook": "Ai hỏi thì trả trong ngày. Chưa đưa bài vào nhóm trường.",
                    "caption": f"Tuan-{n:02d}/README.md", "asset": "—",
                    "cta": "Không đăng thêm", "owner": "Huy",
                    "kpi": week["kpi"], "status": "Việc nội bộ",
                })
            if dow == 2 and n >= 3:
                share = posts_by[(n, 1)]
                rows.append({
                    "date": day, "time": "11:00", "week": n, "channel": "Nhóm học tiếng Nhật",
                    "kind": "Chia sẻ có xin phép", "topic": share["kicker"],
                    "hook": share["hook"],
                    "caption": f"Tuan-{n:02d}/{share['slug']}/caption.txt",
                    "asset": f"Tuan-{n:02d}/{share['slug']}/fb-1080.png",
                    "cta": "Hỏi admin nhóm trước", "owner": "Huy",
                    "kpi": week["kpi"], "status": "Chờ hỏi admin",
                })
            if dow == 4:
                hero = posts_by[(n, 1)]
                need = (MKT / f"Tuan-{n:02d}" / hero["slug"] / "CAN_QUAY.md").exists() or hero.get("footage")
                rows.append({
                    "date": day, "time": "18:00", "week": n, "channel": "TikTok",
                    "kind": "Video dọc 9:16", "topic": hero["kicker"],
                    "hook": hero["hook"],
                    "caption": f"Tuan-{n:02d}/{hero['slug']}/caption.txt",
                    "asset": f"Tuan-{n:02d}/{hero['slug']}/CAN_QUAY.md" if need or hero.get("footage") else f"Tuan-{n:02d}/{hero['slug']}/story-1920.png",
                    "cta": "Xem caption, đừng thêm số", "owner": "Dũng quay · Huy đăng khi có hình",
                    "kpi": week["kpi"],
                    "status": "Cần quay" if hero.get("footage") or hero.get("layout") in ("type", "need") else "Dùng story khi đã quay lời",
                })
            if dow == 6:
                rows.append({
                    "date": day, "time": "20:30", "week": n, "channel": "Nội bộ",
                    "kind": "Đo tuần", "topic": "Điền số thật hoặc «chưa đo»",
                    "hook": "Không ước cho đủ bảng.",
                    "caption": "KeHoach/KeHoach_Marketing_12_Tuan.md", "asset": "—",
                    "cta": "Ba người xem 20 phút", "owner": "Long ghi số · Huy và Dũng đọc",
                    "kpi": week["kpi"], "status": "Việc nội bộ",
                })
            if dow == 6 and n in (5, 8, 12):
                title = {5: "Trailer hoặc teaser", 8: "Demo có thuyết minh", 12: "Nhìn lại 12 tuần"}[n]
                rows.append({
                    "date": day, "time": "21:00", "week": n, "channel": "YouTube",
                    "kind": "Video dài", "topic": title,
                    "hook": title + ". Chỉ đăng khi đã xem hết hình.",
                    "caption": "Video/README.md",
                    "asset": "D:/VTC_Academy/NihongoLife/_work/marketing/",
                    "cta": "Link chơi thật", "owner": "Dũng dựng · Long gắn link · Huy đăng",
                    "kpi": week["kpi"], "status": "Cần quay",
                })
    return rows


def write_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Lich"
    headers = ["Ngày", "Giờ", "Thứ", "Tuần", "Kênh", "Loại bài", "Chủ đề", "Hook", "File caption", "File asset", "CTA", "Người làm", "Mốc tuần", "Trạng thái"]
    header_font = Font(name="Montserrat", bold=True, color="F7BD3D", size=12)
    header_fill = PatternFill("solid", fgColor="111721")
    body_font = Font(name="Montserrat", size=11, color="111721")
    alt = PatternFill("solid", fgColor="FBF6EC")
    thin = Border(
        left=Side(style="thin", color="E4D7B8"), right=Side(style="thin", color="E4D7B8"),
        top=Side(style="thin", color="E4D7B8"), bottom=Side(style="thin", color="E4D7B8"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")
    for col, name in enumerate(headers, 1):
        cell = ws.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
    for r, row in enumerate(calendar_rows(), 2):
        values = [
            row["date"].strftime("%Y-%m-%d"), row.get("time", ""), DOW[row["date"].weekday()], row["week"], row["channel"],
            row["kind"], row["topic"], row["hook"], row["caption"], row["asset"], row["cta"],
            row["owner"], row["kpi"], row["status"],
        ]
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin
            if r % 2 == 0:
                cell.fill = alt
        ws.row_dimensions[r].height = 48
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:N{ws.max_row}"
    widths = [14, 10, 14, 8, 24, 28, 22, 42, 42, 46, 36, 32, 42, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    dv = DataValidation(type="list", formula1='"Sẵn sàng — điền link,Cần quay,Chờ hỏi admin,Việc nội bộ,Đã đăng,Hoãn"', allow_blank=False)
    dv.add(f"N2:N{ws.max_row}")
    ws.add_data_validation(dv)
    _write_year_sheet(wb)
    note = wb.create_sheet("Cach dung", 0)
    note.sheet_view.showGridLines = False
    note["A1"] = "Lịch đăng NihongoLife"
    note["A1"].font = Font(name="Montserrat", bold=True, size=18, color="111721")
    note["A3"] = "Sheet Lich: thứ Hai 26/10/2026 đến chủ Nhật 17/01/2027. Giờ là giờ Việt Nam. Sheet Nam: việc sau 12 tuần, 18/01/2027–16/01/2028, sẽ sửa thường xuyên. Mốc là mục tiêu, không phải kết quả đã có. Chi tiết từng bài Facebook nằm ở chi-tiet.md."
    note["A3"].font = Font(name="Montserrat", size=12, color="111721")
    note["A3"].alignment = Alignment(wrap_text=True)
    note.row_dimensions[3].height = 36
    note.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_year_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Nam")
    headers = ["Ngày", "Quý", "Mùa", "Sự kiện", "Hạn", "Việc", "Người", "Kênh", "File", "Trạng thái"]
    header_font = Font(name="Montserrat", bold=True, color="F7BD3D", size=12)
    header_fill = PatternFill("solid", fgColor="111721")
    body_font = Font(name="Montserrat", size=11, color="111721")
    alt = PatternFill("solid", fgColor="FBF6EC")
    thin = Border(
        left=Side(style="thin", color="E4D7B8"), right=Side(style="thin", color="E4D7B8"),
        top=Side(style="thin", color="E4D7B8"), bottom=Side(style="thin", color="E4D7B8"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")
    for col, name in enumerate(headers, 1):
        cell = ws.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
    for r, row in enumerate(annual_rows(), 2):
        values = [
            row["date"].strftime("%Y-%m-%d"), row["quarter"], row["season"], row["event"],
            row["when"], row["text"], row["owner"], row["channel"], row["folder"], "Chưa làm",
        ]
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin
            if r % 2 == 0:
                cell.fill = alt
        ws.row_dimensions[r].height = 36
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([14, 12, 16, 28, 12, 72, 14, 16, 46, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    dv = DataValidation(type="list", formula1='"Chưa làm,Đang làm,Xong,Hoãn"', allow_blank=False)
    dv.add(f"J2:J{ws.max_row}")
    ws.add_data_validation(dv)


def write_pdf() -> None:
    src = (MKT / "KeHoach" / "KeHoach_Marketing_12_Tuan.md").read_text(encoding="utf-8")
    body = markdown.markdown(src, extensions=["tables", "extra", "sane_lists"])
    cover = """
    <section class="cover">
      <img class="bg" src="file:///workspace/Marketing/HinhAnh/BoiCanh/pho-spawn.png" alt="">
      <div class="veil"></div>
      <div class="inner">
        <div class="mark">NL</div>
        <div class="brand">NIHONGO LIFE</div>
        <div class="jp">日本語ライフ</div>
        <h1>Kế hoạch<br>12 tuần</h1>
        <p class="sub">Học tiếng Nhật bằng cách sống ở Nhật.<br>Bộ này viết cho anh Huy và đội The Next Generation.</p>
        <div class="meta"><span>26.10.2026 — 17.01.2027</span><span>Cộng đồng mở · VTC Academy</span></div>
      </div>
    </section>
    """
    page = f"""<!DOCTYPE html><html lang="vi"><head><meta charset="utf-8">
    <link rel="stylesheet" href="plan-print.css"></head>
    <body>{cover}<main class="doc">{body}</main></body></html>"""
    html_path = TPL / "_plan.html"
    html_path.write_text(page, encoding="utf-8")
    pdf = MKT / "KeHoach" / "KeHoach_Marketing_12_Tuan.pdf"
    if pdf.exists():
        pdf.unlink()
    chrome(
        ["--no-pdf-header-footer", f"--print-to-pdf={pdf}", html_path.as_uri()],
        pdf, timeout=60, min_size=20000,
    )
    print("pdf", pdf.stat().st_size, flush=True)
    write_year_pdf()


def write_year_pdf() -> None:
    src = (MKT / "KeHoachNam" / "KeHoach_Nam_2026-2027.md").read_text(encoding="utf-8")
    body = markdown.markdown(src, extensions=["tables", "extra", "sane_lists"])
    cover = """
    <section class="cover">
      <img class="bg" src="file:///workspace/Marketing/HinhAnh/BoiCanh/pho-ngay.jpg" alt="">
      <div class="veil"></div>
      <div class="inner">
        <div class="mark">NL</div>
        <div class="brand">NIHONGO LIFE</div>
        <div class="jp">日本語ライフ</div>
        <h1>Kế hoạch<br>một năm</h1>
        <p class="sub">Bắt đầu sau 12 tuần. Sẽ được sửa thường xuyên.<br>Viết cho anh Huy và đội The Next Generation.</p>
        <div class="meta"><span>18.01.2027 — 16.01.2028</span><span>Cộng đồng mở · VTC Academy</span></div>
      </div>
    </section>
    """
    page = f"""<!DOCTYPE html><html lang="vi"><head><meta charset="utf-8">
    <link rel="stylesheet" href="plan-print.css"></head>
    <body>{cover}<main class="doc">{body}</main></body></html>"""
    html_path = TPL / "_nam.html"
    html_path.write_text(page, encoding="utf-8")
    pdf = MKT / "KeHoachNam" / "KeHoach_Nam_2026-2027.pdf"
    if pdf.exists():
        pdf.unlink()
    chrome(
        ["--no-pdf-header-footer", f"--print-to-pdf={pdf}", html_path.as_uri()],
        pdf, timeout=60, min_size=20000,
    )
    print("year pdf", pdf.stat().st_size, flush=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--only-pdf", action="store_true")
    parser.add_argument("--only-images", action="store_true")
    parser.add_argument("--skip-images", action="store_true")
    parser.add_argument("--slug", action="append", default=[])
    args = parser.parse_args()
    assert len(POSTS) == 36, len(POSTS)
    assert len(WEEKS) == 12
    build = Path(tempfile.mkdtemp(prefix="nlcards-"))
    if not args.only_pdf:
        write_week_readmes()
        write_posts(build, images=not args.skip_images, slugs=args.slug)
        (MKT / "TuKhoa").mkdir(parents=True, exist_ok=True)
        (MKT / "TuKhoa" / "Ngan_Hang_Tu_Khoa.md").write_text(render_tukhoa(), encoding="utf-8")
        write_tree(MKT / "KeHoachNam")
        xlsx = MKT / "KeHoach" / "Lich_Dang_Bai.xlsx"
        write_xlsx(xlsx)
        bao = ROOT / "Bao_Cao" / "marketing"
        bao.mkdir(parents=True, exist_ok=True)
        shutil.copy(xlsx, bao / "Lich_Noi_Dung_12_Tuan.xlsx")
    if not args.only_images:
        write_pdf()
        pdf = MKT / "KeHoach" / "KeHoach_Marketing_12_Tuan.pdf"
        bao = ROOT / "Bao_Cao" / "marketing"
        bao.mkdir(parents=True, exist_ok=True)
        shutil.copy(pdf, bao / "Ke_Hoach_Marketing_12_Tuan.pdf")
        docs = ROOT / "NihongoLife" / "Docs" / "marketing"
        docs.mkdir(parents=True, exist_ok=True)
        copied = (MKT / "KeHoach" / "KeHoach_Marketing_12_Tuan.md").read_text(encoding="utf-8")
        copied = copied.replace("../HinhAnh/", "../../../Marketing/HinhAnh/")
        (docs / "KE_HOACH_MARKETING_12_TUAN.md").write_text(copied, encoding="utf-8")
    shutil.rmtree(build, ignore_errors=True)


if __name__ == "__main__":
    main()
