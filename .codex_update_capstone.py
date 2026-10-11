from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

src = Path(r"Bao_Cao/NihongoLife_Bao_Cao_Capstone_v4.docx")
tmp = src.with_name(src.stem + ".updated.docx")
doc = Document(src)

def find_paragraph(prefix):
    return next(p for p in doc.paragraphs if p.text.strip().startswith(prefix))

def insert_after(paragraph, text, style="Normal"):
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    out = Paragraph(node, paragraph._parent)
    out.style = style
    out.add_run(text)
    return out

# Chương 10: giữ luận điểm cũ, cập nhật đúng Player HUD vừa nâng cấp.
p = find_paragraph("Đợt hoàn thiện tháng 10 xuất phát")
p.text = (
    "Đợt hoàn thiện tháng 10 xuất phát từ ảnh chụp thực tế cho thấy nhiều lớp phủ cùng lúc: "
    "bảng nhiệm vụ lớn ở góc trái, thanh nút phím tắt cố định ở đáy và các gợi ý chồng nhau. "
    "Giải pháp là một chính sách bố cục dùng chung cho mọi khu vực, trong đó mỗi vùng màn hình "
    "chỉ có một chủ và vùng giữa được giữ thoáng cho quan sát, hội thoại và tương tác."
)
insert_after(p, (
    "Player HUD mới do HUDUI kết nối và StatusDock trình bày: góc trái dưới là portrait lấy trực tiếp "
    "từ model 3D đang điều khiển, tên, địa điểm, cấp độ/EXP, ví và năm vòng Sức khỏe, Thể lực, No, "
    "Khát, Tỉnh táo; góc trái trên chỉ giữ một mục tiêu đang theo dõi; giữa phía trên dành cho zone bar "
    "và HudFeed tối đa ba thông báo; giữa phía dưới chỉ có một prompt tương tác. Card cùng khoá được cập "
    "nhật thay vì xếp chồng, timer tạm dừng khi modal mở và HUD mờ/ẩn trong hội thoại hoặc hoạt động "
    "toàn màn hình. Desktop dùng phím tắt; touch layout giữ action bar, joystick và nút tương tác."
))

# Cập nhật tiêu chí responsive ở bảng yêu cầu phi chức năng.
for table in doc.tables:
    for row in table.rows:
        if row.cells and row.cells[0].text.strip() == "Khả dụng":
            row.cells[1].text = (
                "HUD không chồng lớp ở 1920×1080, 1366×768, 2560×1080, 1024×768, "
                "844×390 và 390×844; mọi cửa sổ đóng được bằng × và Esc"
            )

# Đồng bộ kết luận với ma trận kiểm tra mới.
p = next(p for p in doc.paragraphs if "Chính sách bố cục HUD và điều khiển chuột" in p.text)
p.text = (
    "– Player HUD thống nhất với portrait model thật, năm vòng nhu cầu, một mục tiêu theo dõi, "
    "notification feed chống xếp chồng và điều khiển chuột hai chế độ; thiết kế được kiểm tra theo "
    "ma trận sáu độ phân giải desktop/mobile."
)

# Thêm thông tin đúng của hai video ngay trước Tài liệu tham khảo.
refs = find_paragraph("TÀI LIỆU THAM KHẢO")
heading_node = OxmlElement("w:p")
refs._p.addprevious(heading_node)
heading = Paragraph(heading_node, refs._parent)
heading.style = "Heading 2"
heading.add_run("14.5. Sản phẩm video bàn giao")

video_node = OxmlElement("w:p")
refs._p.addprevious(video_node)
video = Paragraph(video_node, refs._parent)
video.style = "Normal"
video.add_run("Trailer v2: ").bold = True
video.add_run("Bao_Cao/NihongoLife_Trailer_v2.mp4 — H.264/AAC, 1920×1080, 30 fps, 1 phút 25 giây. ")
video.add_run("Gameplay Demo: ").bold = True
video.add_run("Bao_Cao/NihongoLife_Gameplay_Demo.mp4 — H.264/AAC, 1920×1080, 30 fps, 4 phút 18 giây. ")
video.add_run(
    "Hai tệp phục vụ trình chiếu và bàn giao, không được đặt trong Unity Assets và không tham gia runtime."
)

doc.save(tmp)
tmp.replace(src)
print(src)
